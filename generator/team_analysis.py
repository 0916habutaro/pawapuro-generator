"""球団分析：生成球団を球団単位で集計し、実在球団（2022〜2026年版、12球団×5年）と比べる。

- 選手の正規化フレーム（生成選手と実在選手を同じ列にする）
- カテゴリ軸（年齢帯・ポジション・投手役割など）ごとの集計（縦持ち: team_key, axis, group, metric, value, n）
- 主要指標（総合力・投手力・野手力など）と戦力の厚み（ポジション別の何番手）
- 実在の分布の中での位置（百分位・判定）と、実在球団との1対1の比較
- Excel・CSV（zip）の出力

実在の集計ファイル（data/reference/real_team_stats_2022_2026.csv）も、ここの関数で作る（scripts/build_real_team_reference.py）。
生成と実在で同じ集計コードを使い、定義のずれを無くすため。Streamlit には依存しない。
"""
from __future__ import annotations

import csv
import io
import math
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable, Iterable

import numpy as np
import pandas as pd

from generator import real_data
from generator import team as team_lib
from generator.rating import _value as ability_value
from generator.rating import player_rating, ranked_points

APP_DIR = Path(__file__).resolve().parents[1]
SPECIAL_MASTER_PATH = APP_DIR / "data" / "special_abilities.csv"

# ---------------------------------------------------------------------------
# 列・軸・指標の定義
# ---------------------------------------------------------------------------
PITCHER_ABILITIES = ("球速", "コントロール", "スタミナ", "変化球数", "総変化量")
FIELDER_ABILITIES = ("弾道", "ミート", "パワー", "走力", "肩力", "守備力", "捕球")
ABILITY_METRICS = PITCHER_ABILITIES + FIELDER_ABILITIES
SPECIAL_COLUMNS = {"n_blue": "特能_青", "n_red": "特能_赤", "n_gold": "特能_金", "n_green": "特能_緑", "rank_points": "ランク特能点"}
SPECIAL_METRICS = tuple(SPECIAL_COLUMNS.values())

FRAME_COLUMNS = (
    "team_key", "team_label", "name", "uniform_number", "role", "position", "pitcher_role", "throws", "bats", "is_foreign",
    "age", "pro_years", "entry_route", *PITCHER_ABILITIES, *FIELDER_ABILITIES, "rating", *SPECIAL_COLUMNS, "player_class", "archetype",
)
REAL_EXTRA_COLUMNS = ("season", "team", "source", "age_backcalc")

AXIS_ALL = "全体"
AXIS_ROLE = "投手/野手"
AXIS_POSITION = "ポジション"
AXIS_PITCHER_ROLE = "投手役割"
AXIS_AGE = "年齢帯"
AXIS_FOREIGN = "外国人"
AXIS_HAND = "投打"
AXIS_PRO_YEARS = "プロ年数"
AXIS_ENTRY = "入団経路"
AXIS_RATING = "査定帯"
AXIS_CLASS = "選手格"
AXIS_ARCHETYPE = "型"
AXIS_AGE_POSITION = "年齢帯×ポジション"
AXIS_AGE_YEAR = "年齢（1歳刻み）"
AXIS_HEADLINE = "主要指標"
AXIS_DEPTH = "戦力の厚み"
HEADLINE_GROUP = "球団"

POSITION_GROUPS = ("投手", *team_lib.FIELDER_POSITIONS)
AGE_BAND_GROUPS = tuple(team_lib.AGE_BAND_LABELS[band] for band, _low, _high in team_lib.AGE_BANDS)
PRO_YEAR_GROUPS = ("1年目", "2〜3年", "4〜6年", "7〜9年", "10年以上")
ENTRY_ROUTE_GROUPS = ("高卒", "大卒", "社会人", "独立・クラブ", "海外・その他")
# 生成の入団経路 → 実在（2026年版）の区分
ENTRY_ROUTE_MAP = {"海外プロ経由": "海外・その他", "その他": "海外・その他"}
RATING_BAND_GROUPS = ("上位10%", "10〜25%", "25〜50%", "下位50%")
RATING_CUT_LEVELS = (("50%", 0.50), ("75%", 0.75), ("90%", 0.90))
AGE_YEAR_GROUPS = tuple(str(age) for age in range(18, 46))

AXIS_GROUPS: dict[str, tuple[str, ...]] = {
    AXIS_ALL: (AXIS_ALL,),
    AXIS_ROLE: ("投手", "野手"),
    AXIS_POSITION: POSITION_GROUPS,
    AXIS_PITCHER_ROLE: ("先発", "救援"),
    AXIS_AGE: AGE_BAND_GROUPS,
    AXIS_FOREIGN: ("日本人", "外国人"),
    AXIS_HAND: ("右投", "左投", "右打", "左打", "両打"),
    AXIS_PRO_YEARS: PRO_YEAR_GROUPS,
    AXIS_ENTRY: ENTRY_ROUTE_GROUPS,
    AXIS_RATING: RATING_BAND_GROUPS,
    AXIS_AGE_POSITION: tuple(f"{band}×{position}" for band in AGE_BAND_GROUPS for position in POSITION_GROUPS),
    AXIS_AGE_YEAR: AGE_YEAR_GROUPS,
}
# 実在と比べる軸（集計ファイルに入れる軸）
REAL_AXES = (
    AXIS_ALL, AXIS_ROLE, AXIS_POSITION, AXIS_PITCHER_ROLE, AXIS_AGE, AXIS_FOREIGN, AXIS_HAND,
    AXIS_PRO_YEARS, AXIS_ENTRY, AXIS_RATING, AXIS_AGE_POSITION, AXIS_AGE_YEAR,
)
# 生成だけにある軸（実在との比較はしない）
GENERATED_ONLY_AXES = (AXIS_CLASS, AXIS_ARCHETYPE)
ALL_AXES = REAL_AXES + GENERATED_ONLY_AXES
# 画面で選べるカテゴリ軸（交差・1歳刻みは専用の図で出す）
CATEGORY_AXES = (AXIS_ALL, AXIS_ROLE, AXIS_POSITION, AXIS_PITCHER_ROLE, AXIS_AGE, AXIS_FOREIGN, AXIS_HAND, AXIS_PRO_YEARS, AXIS_ENTRY, AXIS_RATING)
# 年齢・プロ年数・入団経路が要る軸。実在の基準は2026年版の12球団だけ
AGE_AXES = (AXIS_AGE, AXIS_PRO_YEARS, AXIS_ENTRY, AXIS_AGE_POSITION, AXIS_AGE_YEAR)
# 人数だけを出す軸（年齢帯×ポジションは平均査定も出す）
COUNT_ONLY_AXES = (AXIS_AGE_POSITION, AXIS_AGE_YEAR)

METRIC_COUNT = "人数"
METRIC_SHARE = "構成比"
METRIC_AGE = "平均年齢"
METRIC_RATING_MEAN = "査定_平均"
METRIC_RATING_MAX = "査定_最大"
# 指標 → 正規化フレームの列（平均を出すもの）
MEAN_METRIC_COLUMNS: dict[str, str] = {
    METRIC_AGE: "age",
    **{name: name for name in ABILITY_METRICS},
    METRIC_RATING_MEAN: "rating",
    **{label: column for column, label in SPECIAL_COLUMNS.items()},
}
GROUP_METRICS = (METRIC_COUNT, METRIC_SHARE, *MEAN_METRIC_COLUMNS, METRIC_RATING_MAX)

# 主要指標（team_rating_metrics と同じ定義）
HEADLINE_METRICS = ("総合力", "投手力", "野手力", "全員平均", "平均年齢", "主力の平均年齢", "外国人数", "左投手数", "先発の厚み", "救援の厚み")
HEADLINE_RATING_METRICS = ("総合力", "投手力", "野手力", "全員平均", "先発の厚み", "救援の厚み")
STARTER_DEPTH_COUNT = 6
RELIEVER_DEPTH_COUNT = 7
AGE_METRICS = (METRIC_AGE, "主力の平均年齢")
# 戦力の厚み: 枠と何番手まで見るか（メインポジションのみ）
DEPTH_SLOTS = (("捕手", 3), ("一塁手", 2), ("二塁手", 2), ("三塁手", 2), ("遊撃手", 2), ("外野手", 5), ("先発", STARTER_DEPTH_COUNT), ("救援", RELIEVER_DEPTH_COUNT))
DEPTH_METRIC = "査定"

JUDGE_OUT = "範囲外"
JUDGE_EDGE = "やや外れ"
JUDGE_IN = "範囲内"
JUDGE_NONE = "比較なし"
# 判定の色（既存の構成チェックと同じ目立たない警告色。範囲外＝赤系、やや外れ＝黄系）
JUDGE_COLORS = {JUDGE_OUT: "#FBE0E0", JUDGE_EDGE: "#FFF1DC"}
VALUE_DECIMALS = 6

GENERATED_PREFIX = "gen:"
UNSAVED_TEAM_KEY = "gen:unsaved"
REAL_PREFIX = "real:"


def real_team_key(season: int, team: str) -> str:
    return f"{REAL_PREFIX}{int(season)}:{team}"


def real_team_label(season: int, team: str) -> str:
    return f"{int(season)} {team}"


def parse_real_team_key(key: str) -> tuple[int, str]:
    _prefix, season, team = key.split(":", 2)
    return int(season), team


# ---------------------------------------------------------------------------
# 特殊能力の分類
# ---------------------------------------------------------------------------
@lru_cache(maxsize=1)
def special_kind_map(path: str = str(SPECIAL_MASTER_PATH)) -> dict[str, str]:
    """特殊能力マスター（data/special_abilities.csv。アプリの master.abilities と同じ）の名前 → 種類。"""
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {str(row["name"]): str(row.get("kind") or "blue") for row in csv.DictReader(f)}


def special_color(name: str, kind: str | None = None) -> str:
    """青・赤・金・緑のどれか。種類はマスターの kind（アプリの special_kind と同じ）。マスターに無い名前は青。"""
    kind = kind or special_kind_map().get(str(name), "blue")
    return {"red": "n_red", "gold": "n_gold", "green": "n_green"}.get(kind, "n_blue")


def special_counts(names: Iterable[str], kinds: Iterable[str | None] | None = None) -> dict[str, int]:
    counts = {"n_blue": 0, "n_red": 0, "n_gold": 0, "n_green": 0}
    names = list(names)
    kinds = list(kinds) if kinds is not None else [None] * len(names)
    for name, kind in zip(names, kinds):
        counts[special_color(name, kind)] += 1
    return counts


# ---------------------------------------------------------------------------
# 正規化フレーム
# ---------------------------------------------------------------------------
def _number(value: Any) -> float:
    parsed = ability_value(value)
    return float(parsed) if parsed is not None else math.nan


def _primary_breaking(balls: Iterable[Any], *, real: bool) -> list[dict[str, Any]]:
    """変化球数・総変化量に数える球（第一球種のみ。compare_real_and_generated_balance.py と同じ数え方）。"""
    out = []
    for ball in balls or []:
        if not isinstance(ball, dict) or ball.get("kind", "breaking") != "breaking":
            continue
        if real and int(ball.get("slot") or 1) > 1:
            continue
        if not real and ball.get("is_second_pitch"):
            continue
        out.append(ball)
    return out


def _movement(ball: dict[str, Any]) -> int:
    return int(ball.get("movement", ball.get("level", 0)) or 0)


def _hand_parts(batting_throwing: Any) -> tuple[str | None, str | None]:
    text = str(batting_throwing or "")
    throws = text[0] if text[:1] in ("右", "左") and text[1:2] == "投" else None
    bats = text[2] if len(text) >= 4 and text[2] in ("右", "左", "両") else None
    return throws, bats


def _position_pitcher_role(player: dict[str, Any]) -> str:
    return str(player.get("position") or "")


def players_frame(
    players: list[dict[str, Any]],
    team_key: str,
    team_label: str,
    pitcher_role_of: Callable[[dict[str, Any]], str] | None = None,
) -> pd.DataFrame:
    """生成選手（dict）の正規化フレーム。

    pitcher_role_of は投手の主役割（先発・中継ぎ・抑え）を返す関数（アプリの primary_pitcher_role を使う）。
    渡さなければ選手の position（先発・中継ぎ・抑え）を使う。
    """
    role_of = pitcher_role_of or _position_pitcher_role
    rows = []
    for p in players:
        abilities = p.get("abilities") if isinstance(p.get("abilities"), dict) else {}
        role = "投手" if p.get("role") == "投手" else "野手"
        throws, bats = _hand_parts(p.get("batting_throwing"))
        row: dict[str, Any] = {
            "team_key": team_key,
            "team_label": team_label,
            "name": str(p.get("name", "")),
            "uniform_number": str(p.get("uniform_number") or ""),
            "role": role,
            "position": "投手" if role == "投手" else str(p.get("position", "")),
            "pitcher_role": ("先発" if role_of(p) == "先発" else "救援") if role == "投手" else None,
            "throws": throws,
            "bats": bats,
            "is_foreign": p.get("roster_origin") == "foreign_import",
            "age": float(p["age"]) if p.get("age") not in (None, "") else math.nan,
            "pro_years": float(p["pro_years"]) if p.get("pro_years") not in (None, "") else math.nan,
            "entry_route": ENTRY_ROUTE_MAP.get(str(p.get("entry_route") or ""), str(p.get("entry_route") or "")) or None,
        }
        for key in ABILITY_METRICS:
            row[key] = math.nan
        if role == "投手":
            for key in ("球速", "コントロール", "スタミナ"):
                row[key] = _number(abilities.get(key))
            primary = _primary_breaking(p.get("breaking_balls") or [], real=False)
            row["変化球数"] = float(len(primary))
            row["総変化量"] = float(sum(_movement(ball) for ball in primary))
        else:
            for key in FIELDER_ABILITIES:
                row[key] = _number(abilities.get(key))
        row["rating"] = float(player_rating(p))
        row.update(special_counts([str(name) for name in p.get("special_abilities") or []]))
        row["rank_points"] = float(ranked_points(abilities.get("ranked_specials") or {}, role))
        row["player_class"] = str(p.get("player_class") or "") or None
        row["archetype"] = str(p.get("archetype") or "") or None
        rows.append(row)
    return _finish_frame(pd.DataFrame(rows, columns=list(FRAME_COLUMNS)))


def real_players_to_frame(rows: list[dict[str, Any]]) -> pd.DataFrame:
    """実在選手（real_data.attach_real_details の結果に season・is_foreign・age・pro_years・entry_route・age_backcalc を
    足したもの）の正規化フレーム。scripts/build_real_team_reference.py とテストで使う。"""
    out = []
    for r in rows:
        season = int(r["season"])
        team = str(r["team"])
        role = "投手" if r.get("role") == "投手" else "野手"
        throws, bats = _hand_parts(r.get("throws_bats"))
        number = r.get("number")
        number_text = "" if number is None or (isinstance(number, float) and math.isnan(number)) or number == "" else str(int(float(number)))
        row: dict[str, Any] = {
            "team_key": real_team_key(season, team),
            "team_label": real_team_label(season, team),
            "name": str(r.get("name", "")),
            "uniform_number": number_text,
            "role": role,
            "position": "投手" if role == "投手" else str(r.get("main_position") or ""),
            "pitcher_role": ("先発" if str(r.get("pitcher_roles") or "")[:1] == "先" else "救援") if role == "投手" else None,
            "throws": throws,
            "bats": bats,
            "is_foreign": bool(r.get("is_foreign")),
            "age": _float_or_nan(r.get("age")),
            "pro_years": _float_or_nan(r.get("pro_years")),
            "entry_route": r.get("entry_route") if isinstance(r.get("entry_route"), str) and r.get("entry_route") else None,
        }
        for key in ABILITY_METRICS:
            row[key] = math.nan
        if role == "投手":
            for key, column in (("球速", "top_speed"), ("コントロール", "control"), ("スタミナ", "stamina")):
                row[key] = _float_or_nan(r.get(column))
            primary = _primary_breaking(r.get("breaking_balls") or [], real=True)
            row["変化球数"] = float(len(primary))
            row["総変化量"] = float(sum(_movement(ball) for ball in primary))
        else:
            for key, column in (("弾道", "trajectory"), ("ミート", "contact"), ("パワー", "power"), ("走力", "run_speed"), ("肩力", "arm_strength"), ("守備力", "fielding"), ("捕球", "catching")):
                row[key] = _float_or_nan(r.get(column))
        rating_dict = real_data.real_player_to_rating_dict(r)
        row["rating"] = float(player_rating(rating_dict))
        names, kinds = [], []
        for name, special_kind in r.get("specials") or []:
            if special_kind in ("rank", "usage"):
                continue
            names.append(name)
            kinds.append("green" if special_kind == "green" else None)
        row.update(special_counts(names, kinds))
        row["rank_points"] = float(ranked_points(rating_dict["abilities"].get("ranked_specials") or {}, role))
        row["player_class"] = None
        row["archetype"] = None
        row["season"] = season
        row["team"] = team
        row["source"] = str(r.get("source") or "")
        row["age_backcalc"] = _float_or_nan(r.get("age_backcalc"))
        out.append(row)
    return _finish_frame(pd.DataFrame(out, columns=list(FRAME_COLUMNS) + list(REAL_EXTRA_COLUMNS)))


def _float_or_nan(value: Any) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return math.nan
    return number


def _finish_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame["uniform_number"] = frame["uniform_number"].fillna("").astype(str)
    frame["is_foreign"] = frame["is_foreign"].astype(bool)
    for column in ("age", "pro_years", *ABILITY_METRICS, "rating", *SPECIAL_COLUMNS):
        frame[column] = pd.to_numeric(frame[column], errors="coerce").astype(float)
    return frame


def real_players_frame(seasons: Iterable[int] | None = None) -> pd.DataFrame | None:
    """実在の選手単位データ（local_data/real_players_2022_2026.csv）。無ければ None。"""
    return real_data.load_real_players(seasons)


# ---------------------------------------------------------------------------
# カテゴリ
# ---------------------------------------------------------------------------
def pro_years_group(years: Any) -> str | None:
    if years is None or pd.isna(years):
        return None
    years = int(years)
    if years <= 1:
        return "1年目"
    if years <= 3:
        return "2〜3年"
    if years <= 6:
        return "4〜6年"
    if years <= 9:
        return "7〜9年"
    return "10年以上"


def rating_cuts_from_frame(frame: pd.DataFrame) -> dict[str, dict[str, float]]:
    """査定帯の境界。実在の全選手の査定の分位点（役割別）。"""
    cuts = {}
    for role in ("投手", "野手"):
        ratings = frame.loc[frame["role"] == role, "rating"].dropna()
        cuts[role] = {label: round(float(ratings.quantile(q)), VALUE_DECIMALS) for label, q in RATING_CUT_LEVELS}
    return cuts


def rating_cuts_from_stats(global_stats: pd.DataFrame) -> dict[str, dict[str, float]]:
    rows = global_stats[global_stats["axis"] == "査定帯の境界"]
    cuts: dict[str, dict[str, float]] = {}
    for row in rows.itertuples():
        cuts.setdefault(str(row.group), {})[str(row.metric)] = float(row.value)
    return cuts


def rating_cut_rows(cuts: dict[str, dict[str, float]]) -> pd.DataFrame:
    rows = [
        {"season": real_data.GLOBAL_SEASON, "team": "全選手", "axis": "査定帯の境界", "group": role, "metric": label, "value": value, "n": 0}
        for role, values in cuts.items() for label, value in values.items()
    ]
    return pd.DataFrame(rows)


def rating_band(rating: float, role: str, cuts: dict[str, dict[str, float]]) -> str | None:
    if rating is None or pd.isna(rating) or role not in cuts:
        return None
    cut = cuts[role]
    if rating >= cut["90%"]:
        return "上位10%"
    if rating >= cut["75%"]:
        return "10〜25%"
    if rating >= cut["50%"]:
        return "25〜50%"
    return "下位50%"


def assign_categories(frame: pd.DataFrame, rating_cuts: dict[str, dict[str, float]]) -> pd.DataFrame:
    """カテゴリ軸の列（列名は軸の名前）を足す。値が決まらない選手は欠損。"""
    out = frame.copy()
    out[AXIS_ALL] = AXIS_ALL
    out[AXIS_ROLE] = out["role"]
    out[AXIS_POSITION] = out["position"].where(out["position"].isin(POSITION_GROUPS))
    out[AXIS_PITCHER_ROLE] = out["pitcher_role"]
    ages = out["age"]
    out[AXIS_AGE] = ages.map(lambda age: team_lib.AGE_BAND_LABELS[team_lib.age_band_of(int(age))] if pd.notna(age) else None)
    out[AXIS_FOREIGN] = np.where(out["is_foreign"], "外国人", "日本人")
    is_pitcher = out["role"] == "投手"
    out[AXIS_HAND] = np.where(
        is_pitcher,
        out["throws"].map(lambda value: f"{value}投" if isinstance(value, str) else None),
        out["bats"].map(lambda value: f"{value}打" if isinstance(value, str) else None),
    )
    out[AXIS_HAND] = out[AXIS_HAND].where(out[AXIS_HAND].isin(AXIS_GROUPS[AXIS_HAND]))
    out[AXIS_PRO_YEARS] = out["pro_years"].map(pro_years_group)
    out[AXIS_ENTRY] = out["entry_route"].where(out["entry_route"].isin(ENTRY_ROUTE_GROUPS))
    out[AXIS_RATING] = [rating_band(rating, role, rating_cuts) for rating, role in zip(out["rating"], out["role"])]
    out[AXIS_CLASS] = out["player_class"]
    out[AXIS_ARCHETYPE] = out["archetype"]
    out[AXIS_AGE_POSITION] = [f"{band}×{position}" if isinstance(band, str) and isinstance(position, str) else None for band, position in zip(out[AXIS_AGE], out[AXIS_POSITION])]
    out[AXIS_AGE_YEAR] = ages.map(lambda age: str(int(age)) if pd.notna(age) else None)
    for axis in ALL_AXES:
        out[axis] = out[axis].astype(object).where(out[axis].notna(), None)
    return out


# ---------------------------------------------------------------------------
# 集計
# ---------------------------------------------------------------------------
def _team_order(frame: pd.DataFrame) -> list[str]:
    return list(dict.fromkeys(frame["team_key"]))


def _expected_groups(axis: str, work: pd.DataFrame) -> list[str]:
    present = [str(value) for value in dict.fromkeys(work[axis].dropna())]
    expected = list(AXIS_GROUPS.get(axis, ()))
    return expected + [value for value in present if value not in expected]


def _share_denominators(axis: str, frame: pd.DataFrame, index: pd.MultiIndex) -> np.ndarray:
    """構成比の分母（球団内の人数。投手役割は投手数、投打は役割ごとの人数）。"""
    is_pitcher = frame["role"] == "投手"
    totals = frame.groupby("team_key").size()
    pitchers = frame[is_pitcher].groupby("team_key").size()
    fielders = frame[~is_pitcher].groupby("team_key").size()
    teams = index.get_level_values("team_key")
    groups = index.get_level_values("group")
    total_values = totals.reindex(teams).fillna(0).to_numpy(dtype=float)
    pitcher_values = pitchers.reindex(teams).fillna(0).to_numpy(dtype=float)
    fielder_values = fielders.reindex(teams).fillna(0).to_numpy(dtype=float)
    if axis == AXIS_PITCHER_ROLE:
        return pitcher_values
    if axis == AXIS_HAND:
        return np.where(groups.str.endswith("投"), pitcher_values, fielder_values)
    return total_values


def group_stats(frame: pd.DataFrame, axis: str) -> pd.DataFrame:
    """カテゴリ軸の group ごとの指標（縦持ち: team_key, axis, group, metric, value, n）。

    frame は assign_categories 済みの正規化フレーム。人数は0人の group も出す（実在の分布に0を含めるため）。
    能力の平均は、その group の投手（投手の能力）・野手（野手の能力）だけで出す。
    """
    columns = ["team_key", "axis", "group", "metric", "value", "n"]
    base = frame
    if axis in AGE_AXES:
        base = frame[frame["age"].notna()]
    work = base[base[axis].notna()]
    teams = _team_order(base)
    if axis in GENERATED_ONLY_AXES:
        teams = _team_order(work)
    if not teams:
        return pd.DataFrame(columns=columns)
    groups = _expected_groups(axis, work)
    index = pd.MultiIndex.from_product([teams, groups], names=["team_key", "group"])
    grouped = work.groupby(["team_key", axis], sort=False)
    counts = grouped.size().rename_axis(["team_key", "group"]).reindex(index, fill_value=0).to_numpy(dtype=float)
    denominators = _share_denominators(axis, base, index)
    with np.errstate(divide="ignore", invalid="ignore"):
        shares = np.where(denominators > 0, counts / denominators * 100, np.nan)
    values = {METRIC_COUNT: counts, METRIC_SHARE: shares}
    valid = {METRIC_COUNT: counts, METRIC_SHARE: counts}
    if axis in COUNT_ONLY_AXES:
        metric_columns = {METRIC_RATING_MEAN: "rating"} if axis == AXIS_AGE_POSITION else {}
    else:
        metric_columns = MEAN_METRIC_COLUMNS
    if metric_columns:
        source_columns = list(dict.fromkeys(metric_columns.values()))
        means = grouped[source_columns].mean().rename_axis(["team_key", "group"]).reindex(index)
        present = grouped[source_columns].count().rename_axis(["team_key", "group"]).reindex(index).fillna(0)
        for metric, column in metric_columns.items():
            values[metric] = means[column].to_numpy(dtype=float)
            valid[metric] = present[column].to_numpy(dtype=float)
        if axis not in COUNT_ONLY_AXES:
            values[METRIC_RATING_MAX] = grouped["rating"].max().rename_axis(["team_key", "group"]).reindex(index).to_numpy(dtype=float)
            valid[METRIC_RATING_MAX] = present["rating"].to_numpy(dtype=float)
    wide = pd.DataFrame(values, index=index)
    n_wide = pd.DataFrame(valid, index=index)
    out = pd.DataFrame({"value": wide.stack(future_stack=True), "n": n_wide.stack(future_stack=True)})
    out.index = out.index.set_names(["team_key", "group", "metric"])
    out = out[out["value"].notna()].reset_index()
    out["value"] = out["value"].round(VALUE_DECIMALS)
    out["n"] = out["n"].astype(int)
    out["axis"] = axis
    return out[columns]


def _mean_or_nan(values: Iterable[float]) -> float:
    values = [float(value) for value in values if value is not None and not pd.isna(value)]
    return float(np.mean(values)) if values else math.nan


def team_headline(frame: pd.DataFrame) -> pd.DataFrame:
    """球団の主要指標（縦持ち: team_key, axis, group, metric, value, n）。査定の定義は team_rating_metrics と同じ。"""
    rows = []
    for team_key in _team_order(frame):
        team = frame[frame["team_key"] == team_key]
        ratings = team["rating"].tolist()
        pitchers = team[team["role"] == "投手"]
        fielders = team[team["role"] != "投手"]
        starters = pitchers[pitchers["pitcher_role"] == "先発"]["rating"].tolist()
        relievers = pitchers[pitchers["pitcher_role"] != "先発"]["rating"].tolist()
        core = team.sort_values("rating", ascending=False, kind="stable").head(team_lib.TOP_TEAM_COUNT)
        values = {
            "総合力": (team_lib._top_mean(ratings, team_lib.TOP_TEAM_COUNT), min(len(ratings), team_lib.TOP_TEAM_COUNT)),
            "投手力": (team_lib._top_mean(pitchers["rating"].tolist(), team_lib.TOP_PITCHER_COUNT), min(len(pitchers), team_lib.TOP_PITCHER_COUNT)),
            "野手力": (team_lib._top_mean(fielders["rating"].tolist(), team_lib.TOP_FIELDER_COUNT), min(len(fielders), team_lib.TOP_FIELDER_COUNT)),
            "全員平均": (float(np.mean(ratings)) if ratings else math.nan, len(ratings)),
            "平均年齢": (_mean_or_nan(team["age"]), int(team["age"].notna().sum())),
            "主力の平均年齢": (_mean_or_nan(core["age"]), int(core["age"].notna().sum())),
            "外国人数": (float(team["is_foreign"].sum()), len(team)),
            "左投手数": (float((pitchers["throws"] == "左").sum()), len(pitchers)),
            "先発の厚み": (team_lib._top_mean(starters, STARTER_DEPTH_COUNT) if starters else math.nan, min(len(starters), STARTER_DEPTH_COUNT)),
            "救援の厚み": (team_lib._top_mean(relievers, RELIEVER_DEPTH_COUNT) if relievers else math.nan, min(len(relievers), RELIEVER_DEPTH_COUNT)),
        }
        for metric in HEADLINE_METRICS:
            value, n = values[metric]
            if value is None or pd.isna(value):
                continue
            rows.append({"team_key": team_key, "axis": AXIS_HEADLINE, "group": HEADLINE_GROUP, "metric": metric, "value": round(float(value), VALUE_DECIMALS), "n": int(n)})
    return pd.DataFrame(rows, columns=["team_key", "axis", "group", "metric", "value", "n"])


def depth_group(slot: str, rank: int) -> str:
    return f"{slot} {rank}番手"


def depth_chart(frame: pd.DataFrame) -> pd.DataFrame:
    """ポジションごとに査定順で並べた何番手の査定と選手名（メインポジションのみ。サブポジションは数えない）。

    その番手の選手がいない枠も、査定を欠損にして出す（手薄な枠）。
    """
    columns = ["team_key", "team_label", "slot", "rank", "rating", "name", "uniform_number", "members"]
    teams = _team_order(frame)
    if not teams:
        return pd.DataFrame(columns=columns)
    work = frame[["team_key", "name", "uniform_number", "rating"]].copy()
    work["slot"] = np.where(frame["role"] == "投手", frame["pitcher_role"], frame["position"])
    work = work.sort_values(["team_key", "rating"], ascending=[True, False], kind="stable")
    work["rank"] = work.groupby(["team_key", "slot"], sort=False).cumcount() + 1
    members = work.groupby(["team_key", "slot"], sort=False).size()
    slots = pd.DataFrame(
        [(team_key, slot, rank) for team_key in teams for slot, count in DEPTH_SLOTS for rank in range(1, count + 1)],
        columns=["team_key", "slot", "rank"],
    )
    out = slots.merge(work, on=["team_key", "slot", "rank"], how="left")
    labels = frame.drop_duplicates("team_key").set_index("team_key")["team_label"]
    out["team_label"] = out["team_key"].map(labels).astype(str)
    out["name"] = out["name"].fillna("").astype(str)
    out["uniform_number"] = out["uniform_number"].fillna("").astype(str)
    out["rating"] = out["rating"].astype(float)
    out["members"] = [int(members.get((team_key, slot), 0)) for team_key, slot in zip(out["team_key"], out["slot"])]
    return out[columns]


def depth_long(depth: pd.DataFrame) -> pd.DataFrame:
    work = depth[depth["rating"].notna()]
    return pd.DataFrame({
        "team_key": work["team_key"],
        "axis": AXIS_DEPTH,
        "group": [depth_group(slot, rank) for slot, rank in zip(work["slot"], work["rank"])],
        "metric": DEPTH_METRIC,
        "value": work["rating"].round(VALUE_DECIMALS),
        "n": 1,
    }).reset_index(drop=True)


def team_long_stats(frame: pd.DataFrame, axes: Iterable[str] = ALL_AXES) -> pd.DataFrame:
    """球団ごとの全集計（カテゴリ軸 → 主要指標 → 戦力の厚み）。実在の集計ファイルもこれで作る。"""
    parts = [group_stats(frame, axis) for axis in axes]
    parts.append(team_headline(frame))
    parts.append(depth_long(depth_chart(frame)))
    return pd.concat([part for part in parts if not part.empty], ignore_index=True)


# ---------------------------------------------------------------------------
# 実在との比較
# ---------------------------------------------------------------------------
def real_long_with_keys(real_stats: pd.DataFrame) -> pd.DataFrame:
    """集計ファイル（season, team, ...）に team_key・team_label を足す。"""
    out = real_stats.copy()
    out["team_key"] = [real_team_key(season, team) for season, team in zip(out["season"], out["team"])]
    out["team_label"] = [real_team_label(season, team) for season, team in zip(out["season"], out["team"])]
    return out


def reference_rows(real: pd.DataFrame, seasons: Iterable[int] | None) -> pd.DataFrame:
    """比較に使う実在の行。年齢系は年版の選択に関わらず2026年版、それ以外は選んだ年版。"""
    seasons = tuple(int(s) for s in (seasons or real_data.DEFAULT_SEASONS))
    age_related = real["axis"].isin(AGE_AXES) | real["metric"].isin(AGE_METRICS)
    keep = (age_related & (real["season"] == real_data.AGE_SEASON)) | (~age_related & real["season"].isin(seasons))
    return real[keep]


def percentile_of(value: float, values: np.ndarray) -> float:
    """実在の中での百分位（0〜100）。同値は中間順位（下回る数＋同値の数の半分）。"""
    if value is None or pd.isna(value) or len(values) == 0:
        return math.nan
    values = np.asarray(values, dtype=float)
    less = float(np.sum(values < value))
    equal = float(np.sum(values == value))
    return (less + 0.5 * equal) / len(values) * 100


def judge(value: float, low: float, p10: float, p90: float, high: float) -> str:
    """実在の最小〜最大の外＝範囲外、10〜90%の外＝やや外れ、それ以外＝範囲内（境界ちょうどは内側）。"""
    if value is None or pd.isna(value) or low is None or pd.isna(low):
        return JUDGE_NONE
    if value < low or value > high:
        return JUDGE_OUT
    if value < p10 or value > p90:
        return JUDGE_EDGE
    return JUDGE_IN


def describe_values(values: Iterable[float]) -> dict[str, float]:
    """最小・10%・中央・90%・最大・平均（分位点は pandas の quantile と同じ線形補間）。"""
    array = np.asarray(list(values) if not isinstance(values, np.ndarray) else values, dtype=float)
    array = array[~np.isnan(array)]
    if array.size == 0:
        return {key: math.nan for key in ("最小", "10%", "中央", "90%", "最大", "平均")}
    p10, median, p90 = np.quantile(array, [0.1, 0.5, 0.9])
    return {
        "最小": float(array.min()), "10%": float(p10), "中央": float(median),
        "90%": float(p90), "最大": float(array.max()), "平均": float(array.mean()),
    }


COMPARE_COLUMNS = (
    "team_key", "team_label", "区分", "グループ", "指標", "値", "人数n",
    "実在_チーム数", "実在_最小", "実在_10%", "実在_中央", "実在_90%", "実在_最大", "実在_平均",
    "百分位", "判定", "実在平均との差", "実在平均との差%",
)
OPPONENT_COLUMNS = ("比較相手", "比較相手の値", "比較相手との差")


def compare_with_real(
    gen: pd.DataFrame,
    real: pd.DataFrame,
    seasons: Iterable[int] | None = None,
    real_team_key: str | None = None,
    team_labels: dict[str, str] | None = None,
) -> pd.DataFrame:
    """生成の集計（team_long_stats）を、実在の集計（real_long_with_keys）の分布と比べる。

    実在の値が無い項目（生成のみの軸など）は判定を「比較なし」にする。
    real_team_key を渡すと、その実在球団の値と差の列を足す（年齢系は2026年版の球団だけ値がある）。
    """
    reference = reference_rows(real, seasons)
    keys = ["axis", "group", "metric"]
    labels = team_labels or {}
    stat_names = ("最小", "10%", "中央", "90%", "最大", "平均")
    stat_rows, pools = [], {}
    for key, values in reference.groupby(keys, sort=False)["value"]:
        array = np.sort(values.to_numpy(dtype=float))
        pools[key] = array
        stats = describe_values(array)
        stat_rows.append((*key, len(array), *(stats[name] for name in stat_names)))
    stat_frame = pd.DataFrame(stat_rows, columns=[*keys, "実在_チーム数", *(f"実在_{name}" for name in stat_names)])
    out = gen.merge(stat_frame, on=keys, how="left")
    out["実在_チーム数"] = out["実在_チーム数"].fillna(0).astype(int)
    values = out["value"].to_numpy(dtype=float)
    percentiles = np.full(len(out), np.nan)
    for key, positions in out.groupby(keys, sort=False).indices.items():
        pool = pools.get(key)
        if pool is None or pool.size == 0:
            continue
        targets = values[positions]
        less = np.searchsorted(pool, targets, side="left")
        less_equal = np.searchsorted(pool, targets, side="right")
        percentiles[positions] = (less + 0.5 * (less_equal - less)) / pool.size * 100
    out["百分位"] = np.where(np.isnan(values), np.nan, percentiles)
    low, p10, p90, high = (out[f"実在_{name}"].to_numpy(dtype=float) for name in ("最小", "10%", "90%", "最大"))
    out["判定"] = np.select(
        [np.isnan(values) | np.isnan(low), (values < low) | (values > high), (values < p10) | (values > p90)],
        [JUDGE_NONE, JUDGE_OUT, JUDGE_EDGE], default=JUDGE_IN,
    )
    mean = out["実在_平均"].to_numpy(dtype=float)
    out["実在平均との差"] = values - mean
    with np.errstate(divide="ignore", invalid="ignore"):
        out["実在平均との差%"] = np.where(mean != 0, (values - mean) / np.abs(mean) * 100, np.nan)
    out["team_label"] = [labels.get(key, key) for key in out["team_key"]]
    out = out.rename(columns={"axis": "区分", "group": "グループ", "metric": "指標", "value": "値", "n": "人数n"})
    out = out[list(COMPARE_COLUMNS)]
    if real_team_key:
        opponent = real[real["team_key"] == real_team_key][keys + ["value"]].rename(
            columns={"axis": "区分", "group": "グループ", "metric": "指標", "value": "比較相手の値"},
        )
        out = out.merge(opponent, on=["区分", "グループ", "指標"], how="left")
        out.insert(len(out.columns) - 1, "比較相手", real_team_label(*parse_real_team_key(real_team_key)))
        out["比較相手との差"] = out["値"] - out["比較相手の値"]
    return out


# ---------------------------------------------------------------------------
# 人数構成（構成チェック表＋百分位）
# ---------------------------------------------------------------------------
def composition_reference(column: str, seasons: Iterable[int] | None = None) -> list[float]:
    """構成項目の実在の値（team.real_composition_values と同じ規則: 抑え・先発の適性は2022〜2025年版、年齢帯は2026年版）。

    seasons で年版を絞る。年齢帯は年版の選択に関わらず2026年版。
    """
    seasons = tuple(int(s) for s in (seasons or real_data.DEFAULT_SEASONS))
    rows = team_lib.load_team_templates()
    if column in ("closer_aptitude", "starter_aptitude"):
        rows = tuple(row for row in rows if row["season"] != 2026)
    if not column.startswith("age_"):
        rows = tuple(row for row in rows if int(row["season"]) in seasons)
    return [float(row[column]) for row in rows if row.get(column) is not None]


def composition_table(players: list[dict[str, Any]], seasons: Iterable[int] | None = None) -> pd.DataFrame:
    """既存の構成チェック表（COMPOSITION_ITEMS・team_composition_counts・実在の範囲）に、百分位と判定を足したもの。"""
    actual = team_lib.team_composition_counts(players)
    rows = []
    for column, label in team_lib.COMPOSITION_ITEMS:
        values = np.asarray(composition_reference(column, seasons), dtype=float)
        stats = describe_values(values)
        value = float(actual.get(column, 0))
        in_range = "○" if len(values) and stats["最小"] <= value <= stats["最大"] else "範囲外"
        rows.append({
            "項目": label, "この球団": int(value),
            "実在の範囲": f"{int(stats['最小'])}〜{int(stats['最大'])}" if len(values) else "",
            "範囲内": in_range,
            "実在_中央": stats["中央"], "実在_10%": stats["10%"], "実在_90%": stats["90%"],
            "百分位": percentile_of(value, values),
            "判定": judge(value, stats["最小"], stats["10%"], stats["90%"], stats["最大"]),
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# まとめ
# ---------------------------------------------------------------------------
@dataclass
class TeamInput:
    """分析する1球団。players は選手 dict（未保存の球団はセッションの dict、保存済みは DB から読み戻した dict）。"""
    team_key: str
    team_label: str
    players: list[dict[str, Any]]
    meta: dict[str, Any] = field(default_factory=dict)


@dataclass
class TeamAnalysis:
    teams: list[TeamInput]
    seasons: tuple[int, ...]
    real_team_key: str | None
    rating_cuts: dict[str, dict[str, float]]
    frame: pd.DataFrame          # 生成選手の正規化フレーム（カテゴリ列付き）
    gen_long: pd.DataFrame       # 生成の集計（縦持ち）
    real_long: pd.DataFrame      # 比較に使った実在の集計（team_key 付き）
    comparison: pd.DataFrame     # compare_with_real の結果
    depth: pd.DataFrame          # 戦力の厚み（生成）
    composition: pd.DataFrame    # 構成チェック表（team_label 付き）

    @property
    def team_labels(self) -> dict[str, str]:
        return {team.team_key: team.team_label for team in self.teams}

    @property
    def real_team_count(self) -> int:
        return int(self.real_long.loc[self.real_long["season"].isin(self.seasons), "team_key"].nunique())

    @property
    def real_age_team_count(self) -> int:
        return int(self.real_long.loc[self.real_long["season"] == real_data.AGE_SEASON, "team_key"].nunique())

    def reference_text(self) -> str:
        return reference_text(self.seasons, self.real_team_count, self.real_age_team_count)


def reference_text(seasons: Iterable[int], team_count: int, age_team_count: int) -> str:
    seasons = sorted(int(s) for s in seasons)
    span = f"{seasons[0]}年版" if len(seasons) == 1 else ("、".join(str(s) for s in seasons) + "年版")
    if len(seasons) > 1 and seasons == list(range(seasons[0], seasons[-1] + 1)):
        span = f"{seasons[0]}〜{seasons[-1]}年版"
    return f"実在の比較基準: {span}の{team_count}チーム（年齢系は2026年版の{age_team_count}チーム）"


def analyze_teams(
    teams: list[TeamInput],
    seasons: Iterable[int] | None = None,
    real_team_key: str | None = None,
    pitcher_role_of: Callable[[dict[str, Any]], str] | None = None,
    real_stats: pd.DataFrame | None = None,
    frames: dict[str, pd.DataFrame] | None = None,
) -> TeamAnalysis:
    """選んだ球団を分析する。frames に球団ごとの正規化フレーム（players_frame の結果）を渡すと作り直さない（キャッシュ用）。"""
    seasons = tuple(sorted(int(s) for s in (seasons or real_data.DEFAULT_SEASONS)))
    stats = real_stats if real_stats is not None else real_data.load_real_team_stats()
    cuts = rating_cuts_from_stats(real_data.load_global_stats()) if real_stats is None else rating_cuts_from_stats(stats[stats["season"] == real_data.GLOBAL_SEASON])
    stats = stats[stats["season"] != real_data.GLOBAL_SEASON]
    real_long = real_long_with_keys(stats)
    parts = []
    for team in teams:
        frame = (frames or {}).get(team.team_key)
        if frame is None:
            frame = players_frame(team.players, team.team_key, team.team_label, pitcher_role_of)
        parts.append(frame)
    frame = assign_categories(pd.concat(parts, ignore_index=True), cuts) if parts else assign_categories(pd.DataFrame(columns=list(FRAME_COLUMNS)), cuts)
    gen_long = team_long_stats(frame)
    labels = {team.team_key: team.team_label for team in teams}
    comparison = compare_with_real(gen_long, real_long, seasons, real_team_key, labels)
    depth = depth_chart(frame)
    composition = pd.concat(
        [composition_table(team.players, seasons).assign(team_label=team.team_label) for team in teams], ignore_index=True,
    ) if teams else pd.DataFrame()
    used = reference_rows(real_long, seasons)
    return TeamAnalysis(teams, seasons, real_team_key, cuts, frame, gen_long, used, comparison, depth, composition)


# ---------------------------------------------------------------------------
# 表示・出力用の表
# ---------------------------------------------------------------------------
def headline_table(analysis: TeamAnalysis) -> pd.DataFrame:
    """主要指標 × 球団（＋実在の分布・百分位・判定）。"""
    rows = analysis.comparison[analysis.comparison["区分"] == AXIS_HEADLINE].copy()
    rows["指標"] = pd.Categorical(rows["指標"], categories=list(HEADLINE_METRICS), ordered=True)
    rows = rows.sort_values(["指標", "team_label"], kind="stable")
    rows["指標"] = rows["指標"].astype(str)
    columns = ["team_label", "指標", "値", "実在_最小", "実在_10%", "実在_中央", "実在_90%", "実在_最大", "実在_平均", "百分位", "判定"]
    columns += [column for column in OPPONENT_COLUMNS if column in rows.columns]
    return rows[columns].reset_index(drop=True)


def headline_pivot(analysis: TeamAnalysis) -> pd.DataFrame:
    """主要指標を球団ごとに並べた表（列＝球団、最後に実在の中央・10〜90%）。"""
    table = headline_table(analysis)
    if table.empty:
        return pd.DataFrame()
    pivot = table.pivot_table(index="指標", columns="team_label", values="値", aggfunc="first", sort=False)
    pivot = pivot.reindex(index=[m for m in HEADLINE_METRICS if m in pivot.index], columns=[t.team_label for t in analysis.teams if t.team_label in pivot.columns])
    real = table.drop_duplicates("指標").set_index("指標")
    pivot["実在_中央"] = real["実在_中央"].reindex(pivot.index)
    pivot["実在_10%"] = real["実在_10%"].reindex(pivot.index)
    pivot["実在_90%"] = real["実在_90%"].reindex(pivot.index)
    return pivot.reset_index()


def flagged_items(analysis: TeamAnalysis, axes: Iterable[str] | None = None) -> pd.DataFrame:
    """範囲外・やや外れの項目（範囲外が先、百分位の極端な順）。"""
    rows = analysis.comparison[analysis.comparison["判定"].isin([JUDGE_OUT, JUDGE_EDGE])].copy()
    if axes is not None:
        rows = rows[rows["区分"].isin(list(axes))]
    rows["_order"] = rows["判定"].map({JUDGE_OUT: 0, JUDGE_EDGE: 1})
    rows["_extreme"] = (rows["百分位"] - 50).abs()
    rows = rows.sort_values(["team_label", "_order", "_extreme"], ascending=[True, True, False], kind="stable")
    return rows.drop(columns=["_order", "_extreme"]).reset_index(drop=True)


def category_table(analysis: TeamAnalysis, axis: str, metrics: Iterable[str]) -> pd.DataFrame:
    """カテゴリ軸 × 指標の比較表（group × 指標。生成値・実在中央・10〜90%・判定）。"""
    metrics = list(metrics)
    rows = analysis.comparison[(analysis.comparison["区分"] == axis) & (analysis.comparison["指標"].isin(metrics))].copy()
    order = {group: index for index, group in enumerate(AXIS_GROUPS.get(axis, ()))}
    rows["_group"] = rows["グループ"].map(lambda group: order.get(group, len(order)))
    rows["_metric"] = rows["指標"].map(metrics.index)
    rows = rows.sort_values(["team_label", "_group", "_metric"], kind="stable").drop(columns=["_group", "_metric"])
    columns = ["team_label", "区分", "グループ", "指標", "値", "人数n", "実在_中央", "実在_10%", "実在_90%", "百分位", "判定"]
    columns += [column for column in OPPONENT_COLUMNS if column in rows.columns]
    return rows[columns].reset_index(drop=True)


def depth_table(analysis: TeamAnalysis) -> pd.DataFrame:
    """戦力の厚み（ポジション×何番手。セルは査定と選手名）に、実在の同じ枠の分布と判定を付ける。"""
    reference = analysis.real_long[analysis.real_long["axis"] == AXIS_DEPTH]
    pools = {group: np.asarray(values, dtype=float) for group, values in reference.groupby("group")["value"]}
    real_team_total = reference["team_key"].nunique()
    rows = []
    for row in analysis.depth.itertuples(index=False):
        group = depth_group(row.slot, row.rank)
        values = pools.get(group, np.asarray([], dtype=float))
        stats = describe_values(values)
        rows.append({
            "team_label": row.team_label, "枠": row.slot, "番手": int(row.rank), "選手": row.name, "背番号": row.uniform_number,
            "査定": row.rating, "実在_10%": stats["10%"], "実在_中央": stats["中央"], "実在_90%": stats["90%"],
            "実在で枠がある割合%": len(values) / real_team_total * 100 if real_team_total else math.nan,
            "判定": judge(row.rating, stats["最小"], stats["10%"], stats["90%"], stats["最大"]) if not pd.isna(row.rating) else "該当者なし",
        })
    return pd.DataFrame(rows)


def depth_findings(analysis: TeamAnalysis) -> dict[str, list[str]]:
    """「手薄なポジション」「層が厚すぎるポジション」の文章。"""
    table = depth_table(analysis)
    findings: dict[str, list[str]] = {"thin": [], "thick": []}
    for (label, slot), group in table.groupby(["team_label", "枠"], sort=False):
        missing = group[group["判定"] == "該当者なし"]
        low = group[(group["査定"] < group["実在_10%"])]
        high = group[(group["査定"] > group["実在_90%"])]
        prefix = f"{label}：" if analysis.teams and len(analysis.teams) > 1 else ""
        if not missing.empty:
            ranks = "・".join(f"{int(r)}番手" for r in missing["番手"])
            share = missing["実在で枠がある割合%"].iloc[0]
            findings["thin"].append(f"{prefix}{slot}の{ranks}がいません（実在で{slot}の{int(missing['番手'].iloc[0])}番手がいる球団は{share:.0f}%）")
        if len(low) >= max(1, math.ceil(len(group) / 2)):
            findings["thin"].append(f"{prefix}{slot}は{len(low)}枠が実在の10%を下回ります（" + "、".join(f"{int(r.番手)}番手 {r.査定:.0f}" for r in low.itertuples()) + "）")
        if len(high) >= max(1, math.ceil(len(group) / 2)):
            findings["thick"].append(f"{prefix}{slot}は{len(high)}枠が実在の90%を上回ります（" + "、".join(f"{int(r.番手)}番手 {r.査定:.0f}" for r in high.itertuples()) + "）")
    return findings


def age_histogram(analysis: TeamAnalysis) -> pd.DataFrame:
    """年齢別（1歳刻み）の人数と、実在（2026年版）の平均人数。"""
    gen = analysis.gen_long[(analysis.gen_long["axis"] == AXIS_AGE_YEAR) & (analysis.gen_long["metric"] == METRIC_COUNT)]
    real = analysis.real_long[(analysis.real_long["axis"] == AXIS_AGE_YEAR) & (analysis.real_long["metric"] == METRIC_COUNT)]
    real_mean = real.groupby("group")["value"].mean()
    labels = analysis.team_labels
    rows = []
    ages = sorted(set(gen["group"]) | set(real["group"]), key=int)
    for team in analysis.teams:
        values = gen[gen["team_key"] == team.team_key].set_index("group")["value"]
        for age in ages:
            rows.append({"team_label": labels[team.team_key], "年齢": int(age), "この球団": float(values.get(age, 0.0)), "実在平均": float(real_mean.get(age, 0.0))})
    return pd.DataFrame(rows)


def age_position_table(analysis: TeamAnalysis) -> pd.DataFrame:
    """年齢帯×ポジションの人数（生成）と、実在（2026年版）の平均人数・差。"""
    gen = analysis.gen_long[(analysis.gen_long["axis"] == AXIS_AGE_POSITION) & (analysis.gen_long["metric"] == METRIC_COUNT)]
    real = analysis.real_long[(analysis.real_long["axis"] == AXIS_AGE_POSITION) & (analysis.real_long["metric"] == METRIC_COUNT)]
    real_mean = real.groupby("group")["value"].mean()
    rows = []
    for team in analysis.teams:
        values = gen[gen["team_key"] == team.team_key].set_index("group")["value"]
        for band in AGE_BAND_GROUPS:
            for position in POSITION_GROUPS:
                group = f"{band}×{position}"
                value = float(values.get(group, 0.0))
                mean = float(real_mean.get(group, math.nan))
                rows.append({"team_label": team.team_label, "年齢帯": band, "ポジション": position, "この球団": value, "実在平均": mean, "差": value - mean})
    return pd.DataFrame(rows)


def player_list_table(analysis: TeamAnalysis) -> pd.DataFrame:
    columns = ["team_label", "uniform_number", "name", "role", "position", "pitcher_role", "throws", "bats", "is_foreign", "age", "pro_years", "entry_route",
               *ABILITY_METRICS, "rating", *SPECIAL_COLUMNS, "player_class", "archetype", AXIS_AGE, AXIS_HAND, AXIS_PRO_YEARS, AXIS_RATING]
    table = analysis.frame[columns].copy()
    table["uniform_number"] = table["uniform_number"].astype(str)
    return table.rename(columns={
        "uniform_number": "背番号", "name": "名前", "role": "役割", "position": "ポジション", "pitcher_role": "投手役割", "throws": "投",
        "bats": "打", "is_foreign": "外国人", "age": "年齢", "pro_years": "プロ年数", "entry_route": "入団経路", "rating": "査定",
        **SPECIAL_COLUMNS, "player_class": "選手格", "archetype": "型", AXIS_AGE: "年齢帯", AXIS_HAND: "投打", AXIS_PRO_YEARS: "プロ年数区分", AXIS_RATING: "査定帯",
    })


def real_reference_table(analysis: TeamAnalysis) -> pd.DataFrame:
    """比較に使った実在の集計値（行＝区分・グループ・指標、列＝年版×球団）。"""
    real = analysis.real_long
    if real.empty:
        return pd.DataFrame()
    pivot = real.pivot_table(index=["axis", "group", "metric"], columns="team_label", values="value", aggfunc="first", sort=False)
    pivot = pivot.reindex(columns=sorted(pivot.columns))
    pivot.index = pivot.index.set_names(["区分", "グループ", "指標"])
    return pivot.reset_index()


def settings_table(analysis: TeamAnalysis, app_version: str, now: Any = None) -> pd.DataFrame:
    now = pd.Timestamp(now if now is not None else pd.Timestamp.now())
    rows = [("出力日時", now.strftime("%Y-%m-%d %H:%M")), ("アプリのバージョン", app_version)]
    for team in analysis.teams:
        meta = team.meta
        parts = [
            f"ID {meta['team_id']}" if meta.get("team_id") else "未保存",
            team.team_label,
            f"seed {meta.get('team_seed', '')}",
            f"戦力レベル {meta.get('strength', '')}",
            f"カラー {meta.get('color', '')}",
        ]
        rows.append(("分析した球団", "／".join(str(part) for part in parts)))
    rows.append(("実在の基準（年版）", "、".join(str(s) for s in analysis.seasons)))
    rows.append(("既定の基準", real_data.DEFAULT_SEASONS_NOTE))
    rows.append(("実在の基準（チーム数）", f"{analysis.real_team_count}チーム"))
    rows.append(("年齢系の基準", f"2026年版の{analysis.real_age_team_count}チームのみ（年齢・年齢帯・プロ年数・入団経路）"))
    rows.append(("外国人数", "実在の外国人数は目安（外国人一覧と登録名で判定）"))
    rows.append(("1対1の比較相手", real_team_label(*parse_real_team_key(analysis.real_team_key)) if analysis.real_team_key else "なし"))
    return pd.DataFrame(rows, columns=["項目", "値"])


EXPORT_SHEETS = ("設定", "概要", "実在比較", "人数構成", "カテゴリ別能力", "年齢", "戦力の厚み", "特殊能力", "選手一覧", "実在基準")


def build_export_tables(analysis: TeamAnalysis, app_version: str = "", now: Any = None) -> dict[str, pd.DataFrame]:
    """出力する表（シート名 → 表）。複数球団のときは各表に team_label 列を持たせる。"""
    comparison = analysis.comparison
    composition = analysis.composition[["team_label", "項目", "この球団", "実在の範囲", "範囲内", "実在_中央", "実在_10%", "実在_90%", "百分位", "判定"]].copy()
    composition.insert(1, "区分", "構成チェック")
    category_counts = comparison[comparison["区分"].isin(CATEGORY_AXES) & (comparison["指標"] == METRIC_COUNT)]
    category_counts = category_counts.rename(columns={"グループ": "項目", "値": "この球団"})[["team_label", "区分", "項目", "この球団", "実在_中央", "実在_10%", "実在_90%", "百分位", "判定"]]
    ability_metrics = [METRIC_AGE, *ABILITY_METRICS, METRIC_RATING_MEAN, METRIC_RATING_MAX]
    abilities = comparison[comparison["区分"].isin(CATEGORY_AXES + GENERATED_ONLY_AXES) & comparison["指標"].isin(ability_metrics)]
    specials = comparison[comparison["区分"].isin(CATEGORY_AXES) & comparison["指標"].isin(SPECIAL_METRICS)]
    detail_columns = ["team_label", "区分", "グループ", "指標", "値", "人数n", "実在_中央", "実在_10%", "実在_90%", "百分位", "判定"]
    detail_columns += [column for column in OPPONENT_COLUMNS if column in comparison.columns]
    ages = pd.concat([
        age_histogram(analysis).assign(表="年齢別人数").rename(columns={"年齢": "項目"}),
        age_position_table(analysis).assign(表="年齢帯×ポジション", 項目=lambda d: d["年齢帯"] + "×" + d["ポジション"]),
    ], ignore_index=True)
    ages = ages[["team_label", "表", "項目", "この球団", "実在平均"]]
    ages["差"] = ages["この球団"] - ages["実在平均"]
    ages["項目"] = ages["項目"].astype(str)
    return {
        "設定": settings_table(analysis, app_version, now),
        "概要": headline_table(analysis),
        "実在比較": comparison.drop(columns=["team_key"]),
        "人数構成": pd.concat([composition, category_counts], ignore_index=True),
        "カテゴリ別能力": abilities[detail_columns].reset_index(drop=True),
        "年齢": ages,
        "戦力の厚み": depth_table(analysis),
        "特殊能力": specials[detail_columns].reset_index(drop=True),
        "選手一覧": player_list_table(analysis),
        "実在基準": real_reference_table(analysis),
    }


def _round_for_export(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    for column in out.columns:
        if pd.api.types.is_float_dtype(out[column]):
            out[column] = out[column].round(2)
    return out


def export_csv_zip_bytes(tables: dict[str, pd.DataFrame]) -> bytes:
    """各表を CSV（UTF-8 BOM）にして zip にまとめる。ファイル名は 01_設定.csv のように番号付き。"""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, (name, table) in enumerate(tables.items(), start=1):
            archive.writestr(f"{index:02d}_{name}.csv", _round_for_export(table).to_csv(index=False).encode("utf-8-sig"))
    return buffer.getvalue()


def _chart_pivot(name: str, table: pd.DataFrame) -> tuple[pd.DataFrame, str, str] | None:
    """グラフ用の小さな表（行＝項目、列＝球団＋実在の中央）とグラフの題・軸名。概要と人数構成だけ。"""
    if table is None or table.empty:
        return None
    if name == "概要":
        data = table[table["指標"].isin(HEADLINE_RATING_METRICS)]
        index, value, order, title, axis = "指標", "値", HEADLINE_RATING_METRICS, "主要指標（査定）と実在の中央", "査定"
    elif name == "人数構成":
        data = table[table["区分"] == AXIS_POSITION]
        index, value, order, title, axis = "項目", "この球団", POSITION_GROUPS, "ポジション別人数と実在の中央", "人数"
    else:
        return None
    if data.empty:
        return None
    pivot = data.pivot_table(index=index, columns="team_label", values=value, aggfunc="first", sort=False)
    pivot = pivot.reindex([item for item in order if item in pivot.index])
    pivot["実在_中央"] = data.drop_duplicates(index).set_index(index)["実在_中央"].reindex(pivot.index)
    pivot.index.name = f"{'ポジション' if name == '人数構成' else index}（グラフ用）"
    return pivot.reset_index(), title, axis


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def _column_widths(table: pd.DataFrame) -> list[float]:
    widths = []
    for column in table.columns:
        sample = [str(column)] + [str(value) for value in table[column].head(200).tolist()]
        width = max(len(text) + sum(1 for ch in text if ord(ch) > 0x2000) for text in sample)
        widths.append(min(max(8, width + 2), 60))
    return widths


# グラフの元の表を置くシート（メインの表の横に並べず、別のシートにまとめる）
CHART_DATA_SHEET = "グラフ用データ"


def _chart_blocks(tables: dict[str, pd.DataFrame]) -> list[tuple[str, pd.DataFrame, str, str, int]]:
    """グラフ用データのシートに縦に並べる表（シート名, 表, 題, 軸名, 見出し行の行番号）。表の間は1行あける。"""
    blocks = []
    row = 1
    for name, table in tables.items():
        chart_block = _chart_pivot(name, _round_for_export(table))
        if chart_block is None:
            continue
        pivot, title, axis_title = chart_block
        blocks.append((name, pivot, title, axis_title, row))
        row += len(pivot) + 2
    return blocks


def export_excel_bytes(tables: dict[str, pd.DataFrame], charts: bool = True) -> bytes:
    """各表を1シートずつ書き、見出しの固定・列幅・判定の色、概要と人数構成に横棒グラフ（Excel のグラフ）を付ける。

    グラフの元の表は最後のシート「グラフ用データ」にまとめる。
    行数が多い（12球団で1万行以上）ので、openpyxl の書き込み専用モードで書く。
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter, quote_sheetname

    book = Workbook(write_only=True)
    bold = Font(bold=True)
    fills = {label: PatternFill(start_color=color.lstrip("#"), end_color=color.lstrip("#"), fill_type="solid") for label, color in JUDGE_COLORS.items()}
    blocks = {block[0]: block for block in _chart_blocks(tables)} if charts else {}

    def header_cell(sheet: Any, value: Any) -> Any:
        cell = WriteOnlyCell(sheet, value=str(value))
        cell.font = bold
        return cell

    for name, raw in tables.items():
        table = _round_for_export(raw)
        sheet = book.create_sheet(title=name)
        for index, width in enumerate(_column_widths(table), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2" if name == "設定" else "B2"
        judge_positions = [index for index, column in enumerate(table.columns) if column in ("判定", "範囲内")]
        sheet.append([header_cell(sheet, column) for column in table.columns])
        for row in table.to_numpy(dtype=object):
            cells = [_cell_value(value) for value in row]
            fill = next((fills[str(row[i])] for i in judge_positions if str(row[i]) in fills), None)
            if fill is not None:
                styled = []
                for value in cells:
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.fill = fill
                    styled.append(cell)
                cells = styled
            sheet.append(cells)
        if name in blocks:
            _block_name, pivot, title, axis_title, top = blocks[name]
            data_sheet = quote_sheetname(CHART_DATA_SHEET)
            last_column = get_column_letter(len(pivot.columns))
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = title
            chart.y_axis.title = axis_title
            chart.height = 9
            chart.width = 18
            chart.add_data(Reference(range_string=f"{data_sheet}!$B${top}:${last_column}${top + len(pivot)}"), titles_from_data=True)
            chart.set_categories(Reference(range_string=f"{data_sheet}!$A${top + 1}:$A${top + len(pivot)}"))
            sheet.add_chart(chart, f"{get_column_letter(len(table.columns) + 2)}2")
    if blocks:
        sheet = book.create_sheet(title=CHART_DATA_SHEET)
        widths: list[float] = []
        for _name, pivot, _title, _axis, _top in blocks.values():
            for index, width in enumerate(_column_widths(pivot)):
                if index < len(widths):
                    widths[index] = max(widths[index], width)
                else:
                    widths.append(width)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for position, (_name, pivot, _title, _axis, _top) in enumerate(blocks.values()):
            if position:
                sheet.append([])
            sheet.append([header_cell(sheet, column) for column in pivot.columns])
            for row in pivot.itertuples(index=False):
                sheet.append([_cell_value(round(value, 1) if isinstance(value, float) else value) for value in row])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def export_file_stem(team_labels: list[str]) -> str:
    """出力ファイル名の球団の部分（1球団なら球団名、複数なら「複数球団」）。"""
    import re

    if len(team_labels) == 1:
        safe = re.sub(r'[\\/:*?"<>|\s]+', "_", str(team_labels[0])).strip("_")
        return safe or "球団"
    return "複数球団"
