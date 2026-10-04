import hashlib
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

TESTS_DIR = Path(__file__).resolve().parent
APP_DIR = TESTS_DIR.parent
sys.path.insert(0, str(APP_DIR))

import app  # noqa: E402
from generator import team as team_lib  # noqa: E402

BASELINE_PATH = TESTS_DIR / "fixtures" / "generate_player_baseline.json"
TEAM_SEEDS = (20261001, 7, 123456789)


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()


# 改修前の app.py で作った指紋と比べるスクリプト。ドラフト候補の野手は文字列のハッシュ順に依存する
# 既存の処理があるため（改修前から）、PYTHONHASHSEED を固定した別プロセスで比べる。
BASELINE_SCRIPT = r"""
import hashlib, json, logging, sys
sys.path.insert(0, sys.argv[1])
import app
baseline = json.load(open(sys.argv[2], encoding="utf-8"))
master = app.load_master_data()
fp = lambda value: hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()
mismatches = []
for key, values in baseline["fingerprints"].items():
    category, role = key.split("|")
    for seed, expected in values.items():
        if fp(app.generate_player(role, category, master, seed=int(seed), used_names=set())) != expected:
            mismatches.append(f"{key}|{seed}")
for seed, expected in baseline["foreign_import_roster"].items():
    if fp(app.generate_foreign_import_roster(int(seed), master)) != expected:
        mismatches.append(f"foreign_import_roster|{seed}")
for seed, expected in baseline["team_roster"].items():
    if fp(app.generate_team_roster(int(seed), master)) != expected:
        mismatches.append(f"team_roster|{seed}")
print(json.dumps(mismatches, ensure_ascii=False))
"""


# 球団生成の結果（選手・TeamProfile・欠番・背番号）の指紋を出すスクリプト。PYTHONHASHSEED を変えて比べる。
TEAM_HASH_SCRIPT = r"""
import hashlib, json, logging, sys
logging.disable(logging.WARNING)
sys.path.insert(0, sys.argv[1])
import app
master = app.load_master_data()
fp = lambda value: hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()
result = {}
for seed in map(int, sys.argv[2:]):
    team = app.generate_team(seed, master=master)
    result[seed] = {
        "players": fp(team["players"]),
        "profile": fp(team["profile"].to_dict()),
        "retired": team["retired_numbers"],
        "numbers": [p["uniform_number"] for p in team["players"]],
        "categories": sorted({p["category"] for p in team["players"]}),
    }
print(json.dumps(result, ensure_ascii=False))
"""


class TeamModeTestBase(unittest.TestCase):
    teams: dict[int, dict] = {}

    @classmethod
    def setUpClass(cls):
        cls.master = app.load_master_data()
        if not TeamModeTestBase.teams:
            TeamModeTestBase.teams = {seed: app.generate_team(seed, master=cls.master) for seed in TEAM_SEEDS}


class GeneratePlayerCompatibilityTest(unittest.TestCase):
    def test_generate_player_without_team_profile_matches_baseline(self):
        baseline = json.loads(BASELINE_PATH.read_text(encoding="utf-8"))
        env = {**os.environ, "PYTHONHASHSEED": baseline["python_hash_seed"], "PYTHONIOENCODING": "utf-8"}
        result = subprocess.run(
            [sys.executable, "-c", BASELINE_SCRIPT, str(APP_DIR), str(BASELINE_PATH)],
            capture_output=True, text=True, encoding="utf-8", env=env, cwd=str(APP_DIR), check=True,
        )
        mismatches = json.loads(result.stdout.strip().splitlines()[-1])
        self.assertEqual(mismatches, [])
        self.assertEqual(sum(len(values) for values in baseline["fingerprints"].values()), 180)

    def test_accept_returns_none_only_when_rejected(self):
        master = app.load_master_data()
        for seed in (1, 2, 3, 4, 5):
            plain = app.generate_player("野手", "架空球団用", master, seed=seed, used_names=set())
            accepted = app.generate_player("野手", "架空球団用", master, seed=seed, used_names=set(), accept=lambda info: True)
            rejected = app.generate_player("野手", "架空球団用", master, seed=seed, used_names=set(), accept=lambda info: False)
            self.assertEqual(fingerprint(plain), fingerprint(accepted))
            self.assertIsNone(rejected)

    def test_accept_stages_and_player_class(self):
        # accept は early（年齢・選手格）→ ポジション・投打 → type（型）の順に呼ばれ、どの段階で断っても None になる
        master = app.load_master_data()
        profile = team_lib.build_team_profile(TEAM_SEEDS[0])
        for seed in (11, 12, 13):
            calls = []

            def record(info):
                calls.append(dict(info))
                return True

            plain = app.generate_player("投手", "架空球団用", master, seed=seed, used_names=set(), team_profile=profile)
            accepted = app.generate_player("投手", "架空球団用", master, seed=seed, used_names=set(), team_profile=profile, accept=record)
            self.assertEqual(fingerprint(plain), fingerprint(accepted))
            self.assertEqual([info.get("stage") for info in calls], ["early", None, "type"])
            # 外国人（foreign_import）になる候補は early の後で選手格を引き直すので、early は国内選手のときだけ比べる
            compared = calls if plain["roster_origin"] == "domestic" else calls[1:]
            self.assertTrue(all(info["player_class"] == plain["player_class"] for info in compared))
            self.assertEqual(calls[2]["archetype"], plain["archetype"])
            for stage in ("early", None, "type"):
                rejected = app.generate_player(
                    "投手", "架空球団用", master, seed=seed, used_names=set(), team_profile=profile,
                    accept=lambda info, stage=stage: info.get("stage") != stage,
                )
                self.assertIsNone(rejected)


class TeamHashSeedTest(unittest.TestCase):
    def test_team_generation_does_not_depend_on_python_hash_seed(self):
        outputs = []
        for hash_seed in ("1", "4242"):
            env = {**os.environ, "PYTHONHASHSEED": hash_seed, "PYTHONIOENCODING": "utf-8"}
            result = subprocess.run(
                [sys.executable, "-c", TEAM_HASH_SCRIPT, str(APP_DIR), *map(str, TEAM_SEEDS)],
                capture_output=True, text=True, encoding="utf-8", env=env, cwd=str(APP_DIR), check=True,
            )
            outputs.append(json.loads(result.stdout.strip().splitlines()[-1]))
        self.assertEqual(outputs[0], outputs[1])
        # 架空球団用（国内）と助っ人外国人用（外国人）の両方を含む球団で確かめている
        for item in outputs[0].values():
            self.assertEqual(item["categories"], ["助っ人外国人用", "架空球団用"])


class TeamTargetsTest(unittest.TestCase):
    def test_targets_stay_in_ranges(self):
        bounds = team_lib.position_bounds()
        for seed in range(300):
            targets = team_lib.build_team_targets(seed)
            self.assertTrue(30 <= targets["pitchers"] <= 39)
            self.assertTrue(28 <= targets["fielders"] <= 38)
            self.assertTrue(63 <= targets["total"] <= 70)
            self.assertEqual(targets["total"], targets["pitchers"] + targets["fielders"])
            self.assertEqual(sum(targets["fielder_targets"].values()), targets["fielders"])
            for position, (low, high) in bounds.items():
                self.assertTrue(low <= targets["fielder_targets"][position] <= high, (seed, position))
            self.assertGreaterEqual(targets["fielder_targets"]["捕手"], 5)
            self.assertGreaterEqual(targets["fielder_targets"]["遊撃手"], 2)
            self.assertGreaterEqual(targets["fielder_targets"]["外野手"], 8)
            pitcher_targets = targets["pitcher_targets"]
            self.assertEqual(pitcher_targets["先発"] + pitcher_targets["救援"], targets["pitchers"])
            self.assertTrue(5 <= pitcher_targets["左投"] <= 15)

    def test_profile_is_deterministic_and_overridable(self):
        first = team_lib.build_team_profile(42)
        self.assertEqual(first, team_lib.build_team_profile(42))
        self.assertIn(first.strength, team_lib.STRENGTH_LABELS)
        fixed = team_lib.build_team_profile(42, color="機動力", color_intensity=1.0, sub_color="")
        self.assertEqual((fixed.color, fixed.color_intensity, fixed.sub_color), ("機動力", 1.0, ""))
        # カラーを固定しても戦力の値は変わらない（乱数は同じ順に全部引く）
        self.assertEqual(fixed.strength_index, first.strength_index)
        self.assertAlmostEqual(fixed.archetype_multipliers["野手"]["俊足"], team_lib.COLOR_EFFECTS["機動力"]["archetype"]["野手"]["俊足"])

    def test_sub_color_is_never_opposite(self):
        for seed in range(2000):
            profile = team_lib.build_team_profile(seed)
            if profile.sub_color:
                self.assertNotEqual(profile.sub_color, profile.color)
                self.assertNotIn(frozenset({profile.color, profile.sub_color}), team_lib.OPPOSITE_COLORS)
                self.assertNotEqual(profile.color, team_lib.NO_COLOR)

    def test_strength_multiplier_formula(self):
        reference = team_lib.STRENGTH_REFERENCE_INDEX
        for label in team_lib.PLAYER_CLASSES:
            self.assertAlmostEqual(team_lib.strength_class_multiplier(label, reference), team_lib.STRONG_CLASS_BASE[label])
            self.assertAlmostEqual(team_lib.strength_class_multiplier(label, -reference), team_lib.WEAK_CLASS_BASE[label])
            self.assertAlmostEqual(team_lib.strength_class_multiplier(label, 0.0), 1.0)

    def test_round_expected_counts_keeps_total_and_mean(self):
        import random

        expected = {"スター級": 1.4, "一軍主力級": 9.7, "一軍控え級": 6.2, "二軍級": 7.9, "若手素材型": 2.5, "ベテラン型": 3.3}
        total = 31
        sums = dict.fromkeys(expected, 0)
        trials = 4000
        for seed in range(trials):
            counts = team_lib.round_expected_counts(expected, total, random.Random(seed))
            self.assertEqual(sum(counts.values()), total)
            for label, value in expected.items():
                # 切り捨てか切り上げのどちらか
                self.assertIn(counts[label], (int(value), int(value) + 1))
                sums[label] += counts[label]
        # 端数は端数に比例する確率で配るので、平均は期待値のまま
        for label, value in expected.items():
            self.assertAlmostEqual(sums[label] / trials, value, delta=0.03)

    def test_assignment_feasible(self):
        self.assertTrue(team_lib.assignment_feasible({"a": 2, "b": 1}, {"x": 1, "y": 2}, {"a": {"x", "y"}, "b": {"x"}}))
        self.assertFalse(team_lib.assignment_feasible({"a": 2, "b": 1}, {"x": 1, "y": 2}, {"a": {"x"}, "b": {"x"}}))
        # 先に入れた区分を付け替えないと全員を割り当てられない場合
        self.assertTrue(team_lib.assignment_feasible({"a": 1, "b": 1, "c": 1}, {"x": 1, "y": 1, "z": 1}, {"a": {"x", "y"}, "b": {"x"}, "c": {"y", "z"}}))
        self.assertFalse(team_lib.assignment_feasible({"a": 1, "b": 1, "c": 1}, {"x": 1, "y": 1, "z": 1}, {"a": {"x", "y"}, "b": {"x"}, "c": {"y"}}))
        self.assertTrue(team_lib.assignment_feasible({"a": 0}, {"x": 0}, {"a": set()}))

    def test_uniform_number_sort_key(self):
        numbers = ["10", "00", "2", "0", "99", "1"]
        self.assertEqual(sorted(numbers, key=team_lib.uniform_number_sort_key), ["0", "00", "1", "2", "10", "99"])

    def test_uniform_age_multiplier_is_continuous(self):
        for band in team_lib.UNIFORM_BANDS:
            self.assertAlmostEqual(team_lib.uniform_age_multiplier(band, team_lib.UNIFORM_AGE_PIVOT), 1.0)
            # 27歳の前後で段差がない
            below = team_lib.uniform_age_multiplier(band, team_lib.UNIFORM_AGE_PIVOT - 0.001)
            above = team_lib.uniform_age_multiplier(band, team_lib.UNIFORM_AGE_PIVOT + 0.001)
            self.assertAlmostEqual(below, above, places=3)
        # 若手は小さい番号を避けて31〜69番へ、ベテランは逆
        self.assertLess(team_lib.uniform_age_multiplier("0-10", 22), 1.0)
        self.assertGreater(team_lib.uniform_age_multiplier("31-69", 22), 1.0)
        self.assertGreater(team_lib.uniform_age_multiplier("0-10", 33), 1.0)
        self.assertLess(team_lib.uniform_age_multiplier("31-69", 33), 1.0)

    def test_foreign_weight_follows_real_foreign_usage(self):
        stats = team_lib.load_uniform_number_stats()
        foreign = {"role": "投手", "position": "先発", "roster_origin": "foreign_import", "age": 27}
        domestic = {**foreign, "roster_origin": "domestic"}
        self.assertEqual(stats["18"]["foreign"], 0)
        # 実在で外国人がいない番号は、外国人にとって国内選手より大幅に重みが小さい
        self.assertLess(
            team_lib.uniform_number_weight("18", foreign, 0.5, stats) / team_lib.uniform_number_weight("18", domestic, 0.5, stats), 0.05,
        )
        # 42番は外国人の重みが国内選手より大きい
        self.assertGreater(team_lib.uniform_number_weight("42", foreign, 0.5, stats), team_lib.uniform_number_weight("42", domestic, 0.5, stats))


class GenerateTeamTest(TeamModeTestBase):
    def test_same_seed_gives_same_team(self):
        again = app.generate_team(TEAM_SEEDS[0], master=self.master)
        original = self.teams[TEAM_SEEDS[0]]
        for key in ("targets", "actual", "retired_numbers", "strength", "color", "sub_color"):
            self.assertEqual(again[key], original[key])
        self.assertEqual(again["profile"], original["profile"])
        self.assertEqual([p["name"] for p in again["players"]], [p["name"] for p in original["players"]])
        self.assertEqual(fingerprint(again["players"]), fingerprint(original["players"]))

    def test_counts_and_targets(self):
        for seed, team in self.teams.items():
            actual = team["actual"]
            with self.subTest(seed=seed):
                self.assertLessEqual(actual["total"], 70)
                self.assertEqual(actual["pitchers"] + actual["fielders"], actual["total"])
                self.assertEqual(len(team["players"]), actual["total"])
                self.assertGreaterEqual(actual["pos_C"], 5)
                self.assertGreaterEqual(actual["pos_SS"], 2)
                if not any(team["relaxed"].values()):
                    self.assertEqual({key: actual[key] for key in team["targets"]}, team["targets"])

    def test_foreign_counts_match_template(self):
        for seed, team in self.teams.items():
            targets = team_lib.build_team_targets(seed)["foreign_targets"]
            with self.subTest(seed=seed):
                self.assertEqual(team["actual"]["foreign_pitchers"], targets["投手"])
                self.assertEqual(team["actual"]["foreign_fielders"], targets["野手"])
                domestic = [p for p in team["players"] if p.get("roster_origin") != "foreign_import"]
                self.assertTrue(all(p["category"] == "架空球団用" and p["roster_origin"] == "domestic" for p in domestic))

    def test_names_are_unique(self):
        for team in self.teams.values():
            names = [p["name"] for p in team["players"]]
            self.assertEqual(len(names), len(set(names)))

    def test_uniform_numbers(self):
        for seed, team in self.teams.items():
            numbers = [p.get("uniform_number") for p in team["players"]]
            with self.subTest(seed=seed):
                self.assertTrue(all(isinstance(number, str) and number for number in numbers))
                self.assertEqual(len(numbers), len(set(numbers)))
                self.assertFalse(set(numbers) & set(team["retired_numbers"]))
                self.assertTrue(set(numbers) <= set(team_lib.UNIFORM_NUMBERS))
                self.assertTrue(all(isinstance(p.get("roster_index"), int) and p.get("team_seed") == seed for p in team["players"]))

    def test_class_and_type_counts_match_targets(self):
        # 球団ごとの散らばり_改修指示.md: 国内選手の選手格の人数（投手・野手別）と、型の人数（野手全員・先発）を目標どおりに作る
        for seed, team in self.teams.items():
            domestic = [p for p in team["players"] if p.get("roster_origin") != "foreign_import"]
            with self.subTest(seed=seed):
                for role in ("投手", "野手"):
                    members = [p for p in domestic if p["role"] == role]
                    self.assertEqual(sum(team["class_targets"][role].values()), len(members))
                    if not team["relaxed"]["選手格の条件"]:
                        self.assertEqual(Counter(p["player_class"] for p in members), Counter({k: v for k, v in team["class_targets"][role].items() if v}))
                fielders = [p for p in domestic if p["role"] == "野手"]
                starters = [p for p in domestic if p["role"] == "投手" and p["position"] == "先発"]
                self.assertEqual(sum(team["type_targets"]["野手"].values()), len(fielders))
                self.assertEqual(sum(team["type_targets"]["先発"].values()), len(starters))
                if not team["relaxed"]["型の条件"]:
                    self.assertEqual(Counter(p["archetype"] for p in fielders), Counter({k: v for k, v in team["type_targets"]["野手"].items() if v}))
                    self.assertEqual(Counter(p["archetype"] for p in starters), Counter({k: v for k, v in team["type_targets"]["先発"].items() if v}))

    def test_class_targets_follow_expected_composition(self):
        # 目標人数の期待値は、選手格の重み（戦力・カラーの倍率と年齢による制限を含む）と年齢帯の目標人数から作る
        profile = team_lib.build_team_profile(TEAM_SEEDS[0])
        age_targets = {"age_u22": 8, "age_23_25": 14, "age_26_29": 18, "age_30_33": 12, "age_34p": 6}
        expected = app.team_class_expected(profile, "投手", 30, age_targets)
        self.assertAlmostEqual(sum(expected.values()), 30)
        self.assertEqual(set(expected), set(team_lib.PLAYER_CLASSES))
        # 23歳以下はベテラン型にならない・27歳以上は若手素材型にならない（choose_player_class の制限）
        young_only = app.team_class_expected(profile, "投手", 10, {"age_u22": 10})
        self.assertEqual(young_only["ベテラン型"], 0.0)
        old_only = app.team_class_expected(profile, "投手", 10, {"age_30_33": 10})
        self.assertEqual(old_only["若手素材型"], 0.0)

    def test_team_profile_only_changes_domestic_weights(self):
        team = self.teams[TEAM_SEEDS[0]]
        foreign = [p for p in team["players"] if p.get("roster_origin") == "foreign_import"]
        self.assertTrue(foreign)
        self.assertTrue(all(p["category"] == "助っ人外国人用" for p in foreign))


class TeamStorageTest(TeamModeTestBase):
    def setUp(self):
        self.original_db_path = app.DB_PATH
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        app.DB_PATH = Path(self.temp_dir.name) / "players.sqlite3"
        app.clear_history_cache()

    def tearDown(self):
        app.DB_PATH = self.original_db_path
        app.clear_history_cache()
        self.temp_dir.cleanup()

    def test_save_team_adds_team_row_and_links_players(self):
        team = self.teams[TEAM_SEEDS[1]]
        app.init_db()
        self.assertEqual(app.next_team_number(), 1)
        team_id = app.save_team(team)
        with sqlite3.connect(app.DB_PATH) as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM teams").fetchone()[0], 1)
            rows = conn.execute("SELECT team_id, roster_index, uniform_number FROM players").fetchall()
        self.assertEqual(len(rows), len(team["players"]))
        self.assertTrue(all(row[0] == team_id for row in rows))
        self.assertEqual(sorted(row[1] for row in rows), list(range(1, len(team["players"]) + 1)))
        row = app.load_team_row(team_id)
        self.assertEqual(row["team_seed"], team["team_seed"])
        self.assertEqual(json.loads(row["retired_numbers_json"]), team["retired_numbers"])
        self.assertEqual(app.next_team_number(), team_id + 1)

    def test_uniform_number_00_survives_database(self):
        team = dict(self.teams[TEAM_SEEDS[2]])
        players = [dict(p) for p in team["players"]]
        if not any(p["uniform_number"] == "00" for p in players):
            # この球団で "00" が空いていれば、1人を "00" に付け替えて読み戻しを確かめる
            players[0]["uniform_number"] = "00"
        team["players"] = players
        app.save_team(team)
        history = app.load_history()
        self.assertIn("00", set(history["uniform_number"]))
        self.assertEqual(sorted(history["uniform_number"]), sorted(p["uniform_number"] for p in players))
        exported = app.team_export_frame(team)
        self.assertIn("00", set(exported["uniform_number"]))
        self.assertEqual(list(exported.columns[:1]), ["team_name"])

    def test_export_excel_has_three_sheets_and_text_numbers(self):
        import io

        import openpyxl

        team = self.teams[TEAM_SEEDS[0]]
        workbook = openpyxl.load_workbook(io.BytesIO(app.team_excel_bytes(team)))
        self.assertEqual(workbook.sheetnames, ["概要", "投手", "野手"])
        for sheet in ("投手", "野手"):
            rows = list(workbook[sheet].iter_rows(values_only=True))
            column = rows[0].index("uniform_number")
            self.assertTrue(all(isinstance(row[column], str) for row in rows[1:]))
        self.assertEqual(workbook["投手"].max_row - 1, team["actual"]["pitchers"])
        self.assertEqual(workbook["野手"].max_row - 1, team["actual"]["fielders"])

    def test_old_database_gets_new_columns(self):
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.execute("CREATE TABLE players (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL DEFAULT '')")
        app.init_db()
        with sqlite3.connect(app.DB_PATH) as conn:
            columns = {row[1] for row in conn.execute("PRAGMA table_info(players)")}
            tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        self.assertTrue({"team_id", "roster_index", "uniform_number"} <= columns)
        self.assertIn("teams", tables)


if __name__ == "__main__":
    unittest.main()
