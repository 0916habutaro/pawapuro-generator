import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

TESTS_DIR = Path(__file__).resolve().parent
APP_DIR = TESTS_DIR.parent
sys.path.insert(0, str(APP_DIR))

import app  # noqa: E402
from generator import real_data  # noqa: E402
from generator import team as team_lib  # noqa: E402
from generator import team_analysis as ta  # noqa: E402

TEAM_SEED = 20261001
REAL_FIXTURE_TEAMS = (("阪神タイガース", 2025), ("中日ドラゴンズ", 2025), ("阪神タイガース", 2026), ("中日ドラゴンズ", 2026))


def fixture_real_rows() -> list[dict]:
    """実在選手の小さなフィクスチャ（attach_real_details の結果と同じ形）。4球団×投手5人・野手5人。"""
    rows = []
    positions = ("捕手", "一塁手", "二塁手", "遊撃手", "外野手")
    for team_index, (team, season) in enumerate(REAL_FIXTURE_TEAMS):
        for i in range(5):
            rows.append({
                "season": season, "team": team, "name": f"投{team_index}{i}", "number": 11 + i, "role": "投手",
                "throws_bats": "左投左打" if i == 0 else "右投右打", "main_position": "投手", "sub_positions": float("nan"),
                "pitcher_roles": "先" if i < 2 else "中抑", "top_speed": 145 + i + team_index, "control": 50 + 3 * i, "stamina": 40 + 5 * i,
                "specials": [("ノビB", "rank"), ("四球", "normal"), ("積極打法", "green"), ("勝利投手", "usage")][: 1 + i % 4],
                "breaking_balls": [{"kind": "breaking", "movement": 3, "slot": 1}, {"kind": "breaking", "movement": 2, "slot": 2}, {"kind": "second_fastball", "movement": 0, "slot": 1}],
                "is_foreign": i == 4, "age": 20 + 3 * i if season == 2026 else float("nan"),
                "pro_years": 1 + 2 * i if season == 2026 else float("nan"), "entry_route": "大卒" if season == 2026 else None,
                "age_backcalc": float("nan"), "source": "fixture",
            })
            rows.append({
                "season": season, "team": team, "name": f"野{team_index}{i}", "number": 0 if i == 0 else 1 + i, "role": "野手",
                "throws_bats": "右投左打" if i % 2 else "右投右打", "main_position": positions[i], "sub_positions": float("nan"),
                "pitcher_roles": float("nan"), "trajectory": 2 + i % 3, "contact": 40 + 2 * i + team_index, "power": 50 + i, "run_speed": 60 - i,
                "arm_strength": 60 + i, "fielding": 50 + i, "catching": 45 + i,
                "specials": [("チャンスC", "rank"), ("広角打法", "normal")][: 1 + i % 2], "breaking_balls": [],
                "is_foreign": False, "age": 22 + 2 * i if season == 2026 else float("nan"),
                "pro_years": 2 + i if season == 2026 else float("nan"), "entry_route": "高卒" if season == 2026 else None,
                "age_backcalc": float("nan"), "source": "fixture",
            })
    return rows


def fixture_real_world() -> tuple[pd.DataFrame, pd.DataFrame]:
    """フィクスチャの実在選手から、構築スクリプトと同じ手順で集計（season, team, axis, ...）を作る。"""
    frame = ta.real_players_to_frame(fixture_real_rows())
    cuts = ta.rating_cuts_from_frame(frame)
    long = ta.team_long_stats(ta.assign_categories(frame, cuts), ta.REAL_AXES)
    keys = [ta.parse_real_team_key(key) for key in long["team_key"]]
    long.insert(0, "season", [season for season, _team in keys])
    long.insert(1, "team", [team for _season, team in keys])
    stats = pd.concat([ta.rating_cut_rows(cuts), long.drop(columns=["team_key"])], ignore_index=True)
    return frame, stats


class TeamAnalysisTestBase(unittest.TestCase):
    team: dict = {}

    @classmethod
    def setUpClass(cls):
        if not TeamAnalysisTestBase.team:
            TeamAnalysisTestBase.team = app.generate_team(TEAM_SEED, team_name="テスト球団", master=app.load_master_data())
        cls.players = TeamAnalysisTestBase.team["players"]
        cls.real_frame, cls.real_stats = fixture_real_world()

    def analyze(self, players=None, **kwargs) -> ta.TeamAnalysis:
        team = ta.TeamInput(ta.UNSAVED_TEAM_KEY, "テスト球団", players or self.players, {"team_seed": TEAM_SEED})
        return ta.analyze_teams([team], pitcher_role_of=app.team_pitcher_role, real_stats=self.real_stats, **kwargs)


class FrameTest(TeamAnalysisTestBase):
    def test_generated_and_real_frames_share_columns(self):
        generated = ta.players_frame(self.players, ta.UNSAVED_TEAM_KEY, "テスト球団", app.team_pitcher_role)
        self.assertEqual(list(generated.columns), list(ta.FRAME_COLUMNS))
        self.assertEqual(list(self.real_frame.columns[: len(ta.FRAME_COLUMNS)]), list(ta.FRAME_COLUMNS))
        self.assertEqual(list(self.real_frame.columns[len(ta.FRAME_COLUMNS):]), list(ta.REAL_EXTRA_COLUMNS))
        for column in ("age", "pro_years", "rating", *ta.ABILITY_METRICS):
            self.assertTrue(pd.api.types.is_float_dtype(generated[column]), column)
            self.assertTrue(pd.api.types.is_float_dtype(self.real_frame[column]), column)
        self.assertEqual(len(generated), len(self.players))
        self.assertTrue(all(isinstance(value, str) for value in generated["uniform_number"]))

    def test_real_frame_counts_primary_breaking_balls_and_specials(self):
        pitcher = self.real_frame[(self.real_frame["role"] == "投手")].iloc[3]
        # 第二球種（slot 2）と第二ストレートは変化球数・総変化量に数えない
        self.assertEqual(pitcher["変化球数"], 1.0)
        self.assertEqual(pitcher["総変化量"], 3.0)
        # 四球（赤）・積極打法（緑）・勝利投手（起用法は数えない）
        self.assertEqual((pitcher["n_red"], pitcher["n_green"], pitcher["n_blue"]), (1.0, 1.0, 0.0))
        self.assertEqual(pitcher["pitcher_role"], "救援")
        self.assertGreater(pitcher["rank_points"], 0)

    def test_usage_like_specials_are_not_counted_in_any_color(self):
        # 調子・投球位置・慎重盗塁・フル出場は実在では起用法の欄にあるので、緑にも青にも数えない（緑特の型_改修指示.md 1-1）
        usage_like = ["投手調子極端", "投手調子安定", "野手調子極端", "野手調子安定", "投球位置左", "投球位置右", "慎重盗塁", "フル出場"]
        self.assertEqual(set(usage_like), set(ta.USAGE_LIKE_SPECIALS))
        for name in usage_like:
            self.assertIsNone(ta.special_color(name))
        counts = ta.special_counts([*usage_like, "テンポ○", "クロスファイヤー", "四球"])
        self.assertEqual(counts, {"n_blue": 1, "n_red": 1, "n_gold": 0, "n_green": 1})
        # マスターに無い名前は、これまでどおり青
        self.assertEqual(ta.special_color("マスターに無い特能"), "n_blue")

    def test_generated_frame_skips_usage_like_specials(self):
        player = dict(next(p for p in self.players if p["role"] == "投手"))
        player["special_abilities"] = ["投手調子極端", "投球位置右", "テンポ○", "クロスファイヤー"]
        row = ta.players_frame([player], ta.UNSAVED_TEAM_KEY, "テスト球団", app.team_pitcher_role).iloc[0]
        self.assertEqual((row["n_green"], row["n_blue"]), (1.0, 1.0))

    def test_pitcher_role_uses_primary_pitcher_role(self):
        generated = ta.players_frame(self.players, ta.UNSAVED_TEAM_KEY, "テスト球団", app.team_pitcher_role)
        pitchers = [p for p in self.players if p["role"] == "投手"]
        expected = ["先発" if app.team_pitcher_role(p) == "先発" else "救援" for p in pitchers]
        self.assertEqual(generated.loc[generated["role"] == "投手", "pitcher_role"].tolist(), expected)


class AggregationTest(TeamAnalysisTestBase):
    def test_group_counts_add_up_to_team_or_role(self):
        analysis = self.analyze()
        counts = analysis.gen_long[analysis.gen_long["metric"] == ta.METRIC_COUNT]
        total = len(self.players)
        pitchers = sum(p["role"] == "投手" for p in self.players)
        for axis in ta.ALL_AXES:
            values = counts[counts["axis"] == axis]
            if axis == ta.AXIS_PITCHER_ROLE:
                self.assertEqual(values["value"].sum(), pitchers, axis)
            elif axis == ta.AXIS_HAND:
                pitch = values[values["group"].str.endswith("投")]["value"].sum()
                bat = values[values["group"].str.endswith("打")]["value"].sum()
                self.assertEqual((pitch, bat), (pitchers, total - pitchers), axis)
            else:
                self.assertEqual(values["value"].sum(), total, axis)
        shares = analysis.gen_long[(analysis.gen_long["axis"] == ta.AXIS_POSITION) & (analysis.gen_long["metric"] == ta.METRIC_SHARE)]
        self.assertAlmostEqual(shares["value"].sum(), 100.0, places=4)

    def test_real_group_counts_add_up(self):
        categorized = ta.assign_categories(self.real_frame, ta.rating_cuts_from_frame(self.real_frame))
        for axis in (ta.AXIS_POSITION, ta.AXIS_FOREIGN, ta.AXIS_RATING):
            stats = ta.group_stats(categorized, axis)
            counts = stats[stats["metric"] == ta.METRIC_COUNT].groupby("team_key")["value"].sum()
            self.assertTrue((counts == 10).all(), axis)
        # 年齢系の軸は年齢が分かる2026年版の球団だけ
        age = ta.group_stats(categorized, ta.AXIS_AGE)
        self.assertEqual(sorted({ta.parse_real_team_key(key)[0] for key in age["team_key"]}), [2026])

    def test_headline_matches_team_rating_metrics(self):
        analysis = self.analyze()
        headline = analysis.gen_long[analysis.gen_long["axis"] == ta.AXIS_HEADLINE].set_index("metric")["value"]
        metrics = team_lib.team_rating_metrics(self.players)
        self.assertAlmostEqual(headline["総合力"], metrics["top28"], places=5)
        self.assertAlmostEqual(headline["投手力"], metrics["pitcher_top"], places=5)
        self.assertAlmostEqual(headline["野手力"], metrics["fielder_top"], places=5)
        self.assertAlmostEqual(headline["全員平均"], metrics["all"], places=5)
        self.assertAlmostEqual(headline["平均年齢"], team_lib.average_age(self.players), places=5)
        self.assertEqual(headline["外国人数"], team_lib.team_composition_counts(self.players)["foreign"])

    def test_depth_chart_lists_every_slot(self):
        depth = ta.depth_chart(ta.players_frame(self.players, "gen:x", "x", app.team_pitcher_role))
        self.assertEqual(len(depth), sum(count for _slot, count in ta.DEPTH_SLOTS))
        catchers = depth[depth["slot"] == "捕手"]["rating"].dropna().tolist()
        self.assertEqual(catchers, sorted(catchers, reverse=True))


class JudgeTest(unittest.TestCase):
    def test_percentile_and_judge_boundaries(self):
        values = np.arange(1, 101, dtype=float)
        stats = ta.describe_values(values)
        self.assertEqual((stats["最小"], stats["最大"]), (1.0, 100.0))
        self.assertAlmostEqual(stats["10%"], 10.9)
        self.assertAlmostEqual(stats["90%"], 90.1)
        judge = lambda value: ta.judge(value, stats["最小"], stats["10%"], stats["90%"], stats["最大"])  # noqa: E731
        self.assertEqual(judge(0.5), ta.JUDGE_OUT)
        self.assertEqual(judge(100.5), ta.JUDGE_OUT)
        self.assertEqual(judge(1.0), ta.JUDGE_EDGE)     # 最小ちょうどは範囲内側（やや外れ）
        self.assertEqual(judge(100.0), ta.JUDGE_EDGE)   # 最大ちょうど
        self.assertEqual(judge(10.9), ta.JUDGE_IN)      # 10%ちょうどは範囲内
        self.assertEqual(judge(90.1), ta.JUDGE_IN)      # 90%ちょうどは範囲内
        self.assertEqual(judge(10.8), ta.JUDGE_EDGE)
        self.assertEqual(judge(float("nan")), ta.JUDGE_NONE)
        # 同値は中間順位
        self.assertEqual(ta.percentile_of(1.0, values), 0.5)
        self.assertEqual(ta.percentile_of(50.0, values), 49.5)
        self.assertEqual(ta.percentile_of(0.0, values), 0.0)
        self.assertEqual(ta.percentile_of(101.0, values), 100.0)
        self.assertEqual(ta.percentile_of(2.0, np.array([1.0, 2.0, 2.0, 3.0])), 50.0)


class RealSelfComparisonTest(TeamAnalysisTestBase):
    def test_real_team_compared_with_itself(self):
        real_long = ta.real_long_with_keys(self.real_stats[self.real_stats["season"] != real_data.GLOBAL_SEASON])
        cuts = ta.rating_cuts_from_stats(self.real_stats[self.real_stats["season"] == real_data.GLOBAL_SEASON])
        key = ta.real_team_key(2026, "阪神タイガース")
        frame = self.real_frame[self.real_frame["team_key"] == key][list(ta.FRAME_COLUMNS)].copy()
        frame["team_key"] = "gen:self"
        gen_long = ta.team_long_stats(ta.assign_categories(frame, cuts))
        comparison = ta.compare_with_real(gen_long, real_long, real_data.REAL_SEASONS, key)
        compared = comparison[comparison["比較相手の値"].notna()]
        self.assertGreater(len(compared), 100)
        self.assertTrue(np.allclose(compared["比較相手との差"], 0.0))
        # 実在の分布の中の百分位は、集計ファイルの値から計算したものと同じ
        row = comparison[(comparison["区分"] == ta.AXIS_HEADLINE) & (comparison["指標"] == "総合力")].iloc[0]
        pool = self.real_stats[(self.real_stats["axis"] == ta.AXIS_HEADLINE) & (self.real_stats["metric"] == "総合力")]["value"].to_numpy()
        self.assertEqual(row["百分位"], ta.percentile_of(row["値"], pool))
        self.assertEqual(row["実在_チーム数"], len(REAL_FIXTURE_TEAMS))
        # 年齢系は2026年版の球団だけが基準
        age = comparison[(comparison["区分"] == ta.AXIS_HEADLINE) & (comparison["指標"] == "平均年齢")].iloc[0]
        self.assertEqual(age["実在_チーム数"], 2)

    def test_season_filter_and_reference_text(self):
        analysis = self.analyze(seasons=[2026])
        self.assertEqual(analysis.real_team_count, 2)
        self.assertEqual(analysis.real_age_team_count, 2)
        self.assertIn("2026年版の2チーム", analysis.reference_text())
        headline = analysis.comparison[(analysis.comparison["区分"] == ta.AXIS_HEADLINE) & (analysis.comparison["指標"] == "総合力")].iloc[0]
        self.assertEqual(headline["実在_チーム数"], 2)

    def test_opponent_columns(self):
        analysis = self.analyze(real_team_key=ta.real_team_key(2025, "中日ドラゴンズ"))
        for column in ta.OPPONENT_COLUMNS:
            self.assertIn(column, analysis.comparison.columns)
        self.assertEqual(set(analysis.comparison["比較相手"]), {"2025 中日ドラゴンズ"})
        rows = analysis.comparison.dropna(subset=["比較相手の値"])
        self.assertTrue(np.allclose(rows["比較相手との差"], rows["値"] - rows["比較相手の値"]))

    def test_generated_only_axes_are_not_compared(self):
        analysis = self.analyze()
        internal = analysis.comparison[analysis.comparison["区分"].isin(ta.GENERATED_ONLY_AXES)]
        self.assertFalse(internal.empty)
        self.assertEqual(set(internal["判定"]), {ta.JUDGE_NONE})


class ExportTest(TeamAnalysisTestBase):
    def test_excel_has_all_sheets_and_keeps_00(self):
        import openpyxl

        players = [dict(p) for p in self.players]
        if not any(p["uniform_number"] == "00" for p in players):
            players[0]["uniform_number"] = "00"
        tables = ta.build_export_tables(self.analyze(players), "test")
        self.assertEqual(list(tables), list(ta.EXPORT_SHEETS))
        workbook = openpyxl.load_workbook(io.BytesIO(ta.export_excel_bytes(tables)))
        self.assertEqual(workbook.sheetnames, [*ta.EXPORT_SHEETS, ta.CHART_DATA_SHEET])
        sheet = workbook["選手一覧"]
        header = [cell.value for cell in sheet[1]]
        column = header.index("背番号") + 1
        numbers = [sheet.cell(row=row, column=column).value for row in range(2, sheet.max_row + 1)]
        self.assertIn("00", numbers)
        self.assertTrue(all(isinstance(number, str) for number in numbers))
        self.assertEqual(len(workbook["概要"]._charts), 1)
        self.assertEqual(len(workbook["人数構成"]._charts), 1)
        self.assertEqual(workbook["実在比較"].freeze_panes, "B2")
        # グラフ用の表はメインの表の横ではなく「グラフ用データ」シートにあり、グラフもそこを参照する
        overview = workbook["概要"]
        self.assertEqual(overview.max_column, len(tables["概要"].columns))
        data = workbook[ta.CHART_DATA_SHEET]
        self.assertEqual(data.cell(row=1, column=1).value, "指標（グラフ用）")
        self.assertIn("ポジション（グラフ用）", [row[0] for row in data.iter_rows(values_only=True)])
        reference = overview._charts[0].series[0].val.numRef.f
        self.assertTrue(reference.startswith(f"'{ta.CHART_DATA_SHEET}'!"), reference)
        settings = dict(workbook["設定"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(settings["既定の基準"], real_data.DEFAULT_SEASONS_NOTE)

    def test_csv_zip_files(self):
        tables = ta.build_export_tables(self.analyze(), "test")
        with zipfile.ZipFile(io.BytesIO(ta.export_csv_zip_bytes(tables))) as archive:
            names = archive.namelist()
            self.assertEqual(names, [f"{index:02d}_{name}.csv" for index, name in enumerate(ta.EXPORT_SHEETS, start=1)])
            self.assertTrue(archive.read("09_選手一覧.csv").startswith("﻿".encode("utf-8")))

    def test_multiple_teams_have_team_label_columns(self):
        teams = [
            ta.TeamInput("gen:a", "球団A", self.players, {}),
            ta.TeamInput("gen:b", "球団B", self.players[:-3], {}),
        ]
        analysis = ta.analyze_teams(teams, pitcher_role_of=app.team_pitcher_role, real_stats=self.real_stats)
        tables = ta.build_export_tables(analysis, "test")
        for name in ("概要", "実在比較", "人数構成", "カテゴリ別能力", "年齢", "戦力の厚み", "特殊能力", "選手一覧"):
            self.assertEqual(set(tables[name]["team_label"]), {"球団A", "球団B"}, name)


class SavedTeamTest(TeamAnalysisTestBase):
    def setUp(self):
        self.original_db_path = app.DB_PATH
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        app.DB_PATH = Path(self.temp_dir.name) / "players.sqlite3"
        app.clear_history_cache()

    def tearDown(self):
        app.DB_PATH = self.original_db_path
        app.clear_history_cache()
        self.temp_dir.cleanup()

    def test_unsaved_and_saved_team_give_same_analysis(self):
        team_id = app.save_team(self.team)
        saved_players = app.load_team_players(team_id)
        self.assertEqual(len(saved_players), len(self.players))
        unsaved = self.analyze()
        saved = self.analyze(saved_players)
        pd.testing.assert_frame_equal(unsaved.comparison, saved.comparison)
        pd.testing.assert_frame_equal(unsaved.composition, saved.composition)
        pd.testing.assert_frame_equal(unsaved.depth, saved.depth)
        listed = app.list_saved_teams()
        self.assertEqual(listed["id"].tolist(), [team_id])


class RealDataFileTest(unittest.TestCase):
    def test_comparison_works_with_stats_file_only(self):
        missing = Path(tempfile.gettempdir()) / "no_such_real_players.csv"
        self.assertIsNone(real_data.load_real_players(path=missing))
        stats = real_data.load_real_team_stats()
        self.assertEqual(stats.groupby("season")["team"].nunique().to_dict(), {2022: 12, 2023: 12, 2024: 12, 2025: 12, 2026: 12})
        cuts = ta.rating_cuts_from_stats(real_data.load_global_stats())
        self.assertEqual(set(cuts), {"投手", "野手"})
        team = TeamAnalysisTestBase.team or app.generate_team(TEAM_SEED, team_name="テスト球団", master=app.load_master_data())
        analysis = ta.analyze_teams([ta.TeamInput("gen:x", "x", team["players"], {})], pitcher_role_of=app.team_pitcher_role)
        # 既定の比較基準は2024〜2026年版（36チーム）
        self.assertEqual(analysis.seasons, (2024, 2025, 2026))
        self.assertEqual(analysis.real_team_count, 36)
        self.assertIn("2024〜2026年版の36チーム", analysis.reference_text())
        everything = ta.analyze_teams([ta.TeamInput("gen:x", "x", team["players"], {})], real_data.REAL_SEASONS, pitcher_role_of=app.team_pitcher_role)
        self.assertEqual(everything.real_team_count, 60)
        self.assertEqual(analysis.real_age_team_count, 12)
        headline = analysis.comparison[analysis.comparison["区分"] == ta.AXIS_HEADLINE]
        self.assertEqual(len(headline), len(ta.HEADLINE_METRICS))
        self.assertFalse(headline["判定"].eq(ta.JUDGE_NONE).any())

    @unittest.skipUnless(real_data.REAL_PLAYERS_PATH.exists(), "実在の選手単位データ（local_data）が無い")
    def test_stats_file_matches_player_file(self):
        """選手単位のファイルから同じ関数で集計し直すと、集計ファイルと一致する（実在と生成で同じ集計コード）。"""
        players = real_data.load_real_players()
        stats = real_data.load_real_team_stats()
        cuts = ta.rating_cuts_from_stats(real_data.load_global_stats())
        key = ta.real_team_key(2025, "阪神タイガース")
        frame = players[players["season"] == 2025][list(ta.FRAME_COLUMNS)]
        frame = frame[frame["team_key"] == key]
        long = ta.team_long_stats(ta.assign_categories(frame, cuts), ta.REAL_AXES)
        expected = stats[(stats["season"] == 2025) & (stats["team"] == "阪神タイガース")]
        merged = long.merge(expected, on=["axis", "group", "metric"], suffixes=("", "_file"))
        self.assertEqual(len(merged), len(long))
        self.assertTrue(np.allclose(merged["value"], merged["value_file"]))


class ValidateTeamModeCompatibilityTest(unittest.TestCase):
    @unittest.skipUnless((APP_DIR / "reports" / "real_powerpro_players_12teams" / "players.csv").exists(), "2026年版の取り込み結果が無い")
    def test_real_team_metrics_unchanged(self):
        """validate_team_mode.py の実在12球団の査定指標（generator/real_data.py へ移す前の値）。

        実在の「対ランナー」を赤特（対ランナー×）として読み直したため、投手の査定が下がった値（対ランナーの取り違えの修正）。
        """
        sys.path.insert(0, str(APP_DIR / "scripts"))
        import validate_team_mode

        metrics = validate_team_mode.real_team_metrics().set_index("team")
        self.assertEqual(len(metrics), 12)
        self.assertAlmostEqual(metrics.loc["オリックスバファローズ", "metric_top28"], 336.857143, places=5)
        self.assertAlmostEqual(metrics.loc["中日ドラゴンズ", "metric_top28"], 328.535714, places=5)
        self.assertAlmostEqual(metrics.loc["北海道日本ハムファイターズ", "metric_pitcher_top"], 368.384615, places=5)
        # 集計ファイルの2026年版の総合力・投手力・野手力も同じ値（同じ変換・同じ査定）
        stats = real_data.load_real_team_stats((2026,))
        headline = stats[stats["axis"] == ta.AXIS_HEADLINE].pivot_table(index="team", columns="metric", values="value")
        for column, metric in (("metric_top28", "総合力"), ("metric_pitcher_top", "投手力"), ("metric_fielder_top", "野手力")):
            self.assertTrue(np.allclose(headline[metric].reindex(metrics.index), metrics[column]), metric)


if __name__ == "__main__":
    unittest.main()
